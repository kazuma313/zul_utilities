from io import BytesIO
from types import SimpleNamespace

import pytest

pytest.importorskip("pptx")
pytest.importorskip("markdown")
pytest.importorskip("xhtml2pdf")

from pptx import Presentation
from pptx.dml.color import RGBColor

from zul.utilities.markdown_converter.md_to_pdf import MarkdownToPDFConverter
from zul.utilities.markdown_converter.md_to_ppt import DynamicMarkdownToPPTXService

DECK = """# Judul Presentasi

---

## Latar Belakang
* Poin utama
  * Poin turunan
** Turunan dengan asterisk
*** Turunan paling dalam
- **Tebal** biasa

---

## Data
| Nama | Nilai |
|------|-------|
| **A** | 1 |
| B | 2 |
"""


@pytest.fixture
def pptx_service():
    return DynamicMarkdownToPPTXService()


def bullets(parsed_slide):
    return [
        (item["text"], item["level"])
        for item in parsed_slide["content"]
        if item["type"] == "bullet"
    ]


# --------------------------------------------------------------------------
# Markdown → PPTX: parsing
# --------------------------------------------------------------------------


def test_first_heading_becomes_slide_title(pptx_service):
    assert pptx_service._parse_slide("# Judul\nteks biasa")["title"] == "Judul"


def test_second_level_heading_is_title_only_when_slide_has_none(pptx_service):
    slide = pptx_service._parse_slide("## Judul Slide\n## Bagian\n### Sub bagian")

    assert slide["title"] == "Judul Slide"
    assert slide["content"] == [
        {"type": "section", "text": "Bagian", "header_level": 2},
        {"type": "section", "text": "Sub bagian", "header_level": 3},
    ]


def test_bullet_level_follows_indentation(pptx_service):
    slide = pptx_service._parse_slide(
        "* satu\n  * dua\n    - tiga\n        * terlalu dalam"
    )

    assert bullets(slide) == [
        ("satu", 0),
        ("dua", 1),
        ("tiga", 2),
        ("terlalu dalam", 2),
    ]


def test_bullet_level_follows_number_of_asterisks(pptx_service):
    slide = pptx_service._parse_slide("* satu\n** dua\n*** tiga")

    assert bullets(slide) == [("satu", 0), ("dua", 1), ("tiga", 2)]


def test_bold_text_inside_bullet_is_not_treated_as_nesting(pptx_service):
    slide = pptx_service._parse_slide("* **Penting** sekali")

    assert bullets(slide) == [("Penting sekali", 0)]


def test_paragraph_starting_with_bold_is_plain_text(pptx_service):
    slide = pptx_service._parse_slide("**Catatan** untuk pembaca")

    assert slide["content"] == [{"type": "text", "text": "**Catatan** untuk pembaca"}]


def test_table_is_parsed_without_separator_row_and_markup(pptx_service):
    slide = pptx_service._parse_slide(
        "## Data\n| Nama | Nilai |\n|---|---|\n| **A** | 1 |"
    )

    assert slide["table_data"] == [["Nama", "Nilai"], ["A", "1"]]
    assert slide["content"] is None


def test_image_reference_is_collected(pptx_service):
    slide = pptx_service._parse_slide("## Grafik\n![hasil](charts/hasil.png)")

    assert slide["image_paths"] == ["charts/hasil.png"]


# --------------------------------------------------------------------------
# Markdown → PPTX: output
# --------------------------------------------------------------------------


def test_convert_to_bytes_creates_one_slide_per_section(pptx_service):
    deck = Presentation(BytesIO(pptx_service.convert_to_bytes(DECK)))

    assert len(deck.slides) == 3


def test_convert_to_bytes_starts_from_a_fresh_presentation_each_time(pptx_service):
    pptx_service.convert_to_bytes(DECK)

    deck = Presentation(BytesIO(pptx_service.convert_to_bytes(DECK)))

    assert len(deck.slides) == 3


def test_convert_markdown_saves_file(pptx_service, tmp_path):
    output = tmp_path / "deck.pptx"

    pptx_service.convert_markdown(DECK, str(output))

    assert len(Presentation(str(output)).slides) == 3


def test_old_pptx_import_path_still_works():
    from zul.utilities.md_to_ppt import DynamicMarkdownToPPTXService as OldPath

    assert OldPath is DynamicMarkdownToPPTXService


# --------------------------------------------------------------------------
# Markdown → PDF
# --------------------------------------------------------------------------


def test_default_styles_use_page_size_and_margin():
    styles = MarkdownToPDFConverter(page_size="Letter", margin="1.5cm").get_styles()

    assert "size: Letter;" in styles
    assert "margin: 1.5cm;" in styles


def test_plain_custom_css_is_used_as_is():
    converter = MarkdownToPDFConverter()
    converter.set_custom_styles("body { color: red; }")

    assert converter.get_styles() == "body { color: red; }"


def test_custom_css_template_still_gets_page_settings():
    converter = MarkdownToPDFConverter(page_size="A5")
    converter.set_custom_styles("@page {{ size: {page_size}; }}")

    assert converter.get_styles() == "@page { size: A5; }"


def test_custom_styles_do_not_leak_to_other_converters():
    MarkdownToPDFConverter().set_custom_styles("body { color: red; }")

    assert "size: A4;" in MarkdownToPDFConverter().get_styles()


def test_markdown_table_becomes_html_table():
    html = MarkdownToPDFConverter().markdown_to_html("| A | B |\n|---|---|\n| 1 | 2 |")

    assert "<table>" in html


def test_convert_to_bytes_returns_pdf():
    pdf = MarkdownToPDFConverter().convert_to_bytes("# Judul\n\nIsi **dokumen**.")

    assert pdf.startswith(b"%PDF")


def test_convert_writes_pdf_into_output_directory(tmp_path):
    output_dir = tmp_path / "hasil"

    succeeded = MarkdownToPDFConverter().convert(
        "# Judul", str(output_dir), "dokumen.pdf"
    )

    assert succeeded
    assert (output_dir / "dokumen.pdf").read_bytes().startswith(b"%PDF")


def test_pdf_errors_reported_by_xhtml2pdf_give_false_and_empty_bytes(
    monkeypatch, tmp_path
):
    from zul.adapters import xhtml2pdf as xhtml2pdf_adapter

    monkeypatch.setattr(
        xhtml2pdf_adapter.pisa, "CreatePDF", lambda html, dest: SimpleNamespace(err=1)
    )
    converter = MarkdownToPDFConverter()

    assert converter.convert("# Judul", str(tmp_path), "dokumen.pdf") is False
    assert converter.convert_to_bytes("# Judul") == b""


# --------------------------------------------------------------------------
# Markdown → PPTX: warna
# --------------------------------------------------------------------------


def title_color(pptx_bytes, title):
    slide = Presentation(BytesIO(pptx_bytes)).slides[0]
    for shape in slide.shapes:
        for paragraph in shape.text_frame.paragraphs:
            if paragraph.text == title:
                return paragraph.font.color.rgb
    raise AssertionError(f"judul {title!r} tidak ditemukan")


def test_default_style_colors_are_still_pptx_rgbcolor(pptx_service):
    style = pptx_service.get_default_style()

    assert isinstance(style["h2"]["font_color"], RGBColor)
    assert style["h2"]["font_color"] == RGBColor(230, 126, 34)
    assert style["table"]["header_color"] == RGBColor(0, 51, 102)


@pytest.mark.parametrize("color", [(26, 54, 93), "#1A365D", RGBColor(26, 54, 93)])
def test_title_color_can_be_tuple_hex_or_rgbcolor(color):
    service = DynamicMarkdownToPPTXService(style_config={"h2": {"font_color": color}})

    assert title_color(service.convert_to_bytes("## Judul"), "Judul") == RGBColor(
        26, 54, 93
    )


def test_default_title_color_is_applied(pptx_service):
    pptx_bytes = pptx_service.convert_to_bytes("## Judul")

    assert title_color(pptx_bytes, "Judul") == RGBColor(230, 126, 34)
