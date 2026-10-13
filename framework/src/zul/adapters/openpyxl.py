"""
Adapter openpyxl: membaca baris file Excel .xlsx.

Gunanya:
    Satu-satunya file Zul yang mengimpor openpyxl. Baris pertama sheet
    dibaca sebagai nama kolom, dan setiap baris berikutnya menjadi satu dict.
    Butuh extra evaluation: `pip install "zul[evaluation]"`.

Cara pakai:
    from zul.adapters import openpyxl as excel

    rows = excel.read_rows("dataset.xlsx")              # sheet pertama
    rows = excel.read_rows("dataset.xlsx", sheet="qa")  # sheet tertentu
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def read_rows(path: str | Path, sheet: str | None = None) -> list[dict[str, Any]]:
    """Setiap baris sheet sebagai dict {nama kolom: nilai}; baris kosong dilewati.

    Kolom tanpa judul dinamai `kolom_N`, dengan N mulai dari 1. Sel berisi
    rumus dibaca sebagai nilai hasil terakhir yang tersimpan di file.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []
        names = [
            str(name).strip() if name is not None else f"kolom_{i}"
            for i, name in enumerate(header, 1)
        ]
        return [
            dict(zip(names, values, strict=False))
            for values in rows
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()
