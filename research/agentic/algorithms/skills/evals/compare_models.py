"""Run every skill of this pool with several models on the SAME requests and build one comparison page.

    python evals/compare_models.py gemma3:4b qwen3:8b          # local models through Ollama
    python evals/compare_models.py --claude                    # the specs in evals/comparison/claude/ (written by Claude)
    python evals/compare_models.py --report                    # rebuild evals/comparison/index.html only
    python evals/compare_models.py gemma3:4b --replay          # after a skill was repaired: build again from the
                                                               # answers saved in the last run, without asking the model

Each model gets the same material and the same request per skill.  What is kept for every case:
the arguments or spec the model wrote, the file the skill built from it, a preview, and the time.
evals/comparison/notes.json holds what a reader should notice per case and model; it is written by
whoever reviewed the results, not by this script.

  tool cases     the model writes the arguments of one tool call; the tool is run with them
                 (web search, image and subagents are not run: they need the network or an upload)
  runner cases   the model writes the spec through generate_deck.py / generate_poster.py /
                 generate_mindmap.py; the skill's own script builds the file

"Claude" has no local server: its arguments and specs are files in evals/comparison/claude/, and the same
tools and build scripts are run with them.  Previews: PowerPoint (COM) exports the slides as PNG;
Word and Excel files are read back with python-docx and openpyxl.
"""

from __future__ import annotations

import argparse
import html
import json
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.error
from pathlib import Path

HERE = Path(__file__).resolve().parent
POOL = HERE.parent
OUT = HERE / "comparison"
sys.path.insert(0, str(HERE))
import small_model_check as smc  # noqa: E402

MATERIAL = """Survei kebiasaan belanja online mahasiswa (data contoh).
Responden: 350 mahasiswa dari 500 yang diundang, usia 18-24 tahun, 5 kampus di Medan, Maret-April 2026.
Metode: survei online 24 pertanyaan, 12 wawancara mendalam, dan analisis riwayat transaksi.
Temuan: 72% belanja online minimal sebulan sekali. Alasan utama: harga 54%, kemudahan 31%, pilihan produk 15%. Rata-rata belanja Rp 310.000 per bulan. 64% membayar dengan e-wallet. 41% pernah menyesal membeli karena diskon.
Tren: belanja lewat live streaming naik dari 12% (2024) menjadi 29% (2026). 23% memakai paylater.
Rekomendasi: mahasiswa membuat anggaran bulanan; kampus mengadakan kelas literasi keuangan; platform menampilkan batas paylater dengan jelas.
"""

# case id -> (skill folder, tool name, title, request, run the tool?)
TOOL_CASES = {
    "calculator": ("calculator", "math_calculator", "Kalkulator",
                   "Dari 500 mahasiswa yang diundang, 350 mengisi survei. Berapa persen tingkat responsnya?", True),
    "chart": ("chart", "generate_chart", "Grafik",
              "Buat grafik batang alasan utama belanja online: harga 54%, kemudahan 31%, pilihan produk 15%.", True),
    "docx": ("docx", "create_docx", "Dokumen Word",
             "Buat ringkasan eksekutif survei berikut sebagai dokumen Word: satu paragraf pembuka, tiga butir temuan utama, "
             "dan tabel alasan utama.\n\n" + MATERIAL, True),
    "xlsx": ("xlsx", "create_xlsx", "Spreadsheet Excel",
             "Buat file Excel dari survei berikut dengan dua sheet: 'Alasan' (alasan, persen) dan 'Tren' (tahun, persen "
             "belanja lewat live streaming).\n\n" + MATERIAL, True),
    "mind_map_tool": ("mind_map", "create_mind_map", "Mind map (tool)",
                      "Buat mind map dari ringkasan survei berikut, simpan sebagai survei.svg.\n\n" + MATERIAL, True),
    "web_search": ("web_search", "web_search", "Web search (argumen saja)",
                   "Cari data terbaru jumlah pengguna e-commerce di Indonesia.", False),
    "image": ("image", "view_image", "Image reader (argumen saja)",
              "Jelaskan grafik pada gambar dengan file_id t1/alasan.png. Batang mana yang paling tinggi?", False),
    "subagents": ("subagent_research", "dispatch_research_subagents", "Subagent research (argumen saja)",
                  "Riset tiga hal secara paralel untuk melengkapi survei ini: tren live shopping di Indonesia, pemakaian "
                  "paylater oleh mahasiswa, dan aturan perlindungan konsumen e-commerce.", False),
}

# case id -> (skill folder, title, generator script, topic, output file, build script for a ready spec, spec file name)
RUNNER_CASES = {
    "mind_map": ("mind_map", "Mind map (dari dokumen)", "scripts/generate_mindmap.py", None, "map.svg",
                 "scripts/build_mindmap.py", "mind_map.md"),
    "poster": ("resaerch_poster", "Poster penelitian", "scripts/generate_poster.py",
               "Poster hasil survei kebiasaan belanja online mahasiswa", "poster.png", "scripts/build_poster.py", "poster.json"),
    "pptx_claude": ("pptx_claude", "Slide presentasi (pptx_claude)", "scripts/generate_deck.py",
                    "Presentasi hasil survei kebiasaan belanja online mahasiswa", "deck.pptx", "scripts/create_pptx.py",
                    "pptx_claude.json"),
    "pptx_research": ("pptx_research", "Slide proposal penelitian (pptx_research)", "scripts/generate_deck.py",
                      "Proposal penelitian: pengaruh paylater terhadap kebiasaan belanja mahasiswa", "proposal.pptx",
                      "scripts/build_deck.py", "pptx_research.json"),
}
ORDER = ["mind_map", "mind_map_tool", "poster", "pptx_claude", "pptx_research", "docx", "xlsx", "chart", "calculator",
         "web_search", "image", "subagents"]
LABELS = {"claude": "Claude (sesi ini)"}
FOLDERS = re.compile(r"[A-Za-z]:[\\/](?:[^\\/\s]+[\\/])+")       # the folders of an absolute Windows path; the page shows file names only


def slug(model: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", model.lower()).strip("-")


def new_files(folder: Path, before: set) -> list[str]:
    return sorted(p.name for p in folder.iterdir() if p.name not in before and p.is_file())


# ---------------------------------------------------------------------------
# Tool cases
# ---------------------------------------------------------------------------

def model_arguments(args, row: dict, tool, request: str) -> dict:
    """One tool call from the model: its own tool calling, or JSON held to the tool's schema."""
    from langchain_core.utils.function_calling import convert_to_openai_tool

    spec = convert_to_openai_tool(tool)
    body = row["skill_md"].split("---", 2)[-1].strip()
    user = {"role": "user", "content": request}
    if args.calls == "native":
        reply = smc.chat(args, [{"role": "system", "content": "You have one skill. Use its tool to do what the user asks.\n\n" + body}, user],
                         tools=[spec], max_tokens=args.max_tokens)
        calls = reply.get("tool_calls") or []
        if not calls:
            raise ValueError("no tool call; the model answered: " + " ".join((reply.get("content") or "").split())[:120])
        call_args = calls[0].get("function", {}).get("arguments") or {}
    else:
        system = (f"You have one skill. Do what the user asks by calling its tool `{tool.name}`.\n"
                  f"Reply with the arguments of that call as one JSON object, nothing else.\n\n{body}\n\n"
                  f"Tool `{tool.name}`:\n{spec['function'].get('description', '')}")
        reply = smc.chat(args, [{"role": "system", "content": system}, user], max_tokens=args.max_tokens,
                         schema=spec["function"]["parameters"])
        call_args = reply.get("content") or ""
    return json.loads(call_args) if isinstance(call_args, str) else call_args


def run_tool_case(case: str, rows: dict, folder: Path, get_arguments) -> dict:
    skill, tool_name, _, request, execute = TOOL_CASES[case]
    record = {"case": case, "ok": False, "seconds": 0.0, "arguments": None, "result": "", "files": []}
    tool = rows[skill]["tools"].get(tool_name)
    if tool is None:
        record["result"] = "skill tidak bisa dimuat: " + (rows[skill]["error"] or "tool tidak ditemukan")
        return record
    folder.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    try:
        record["arguments"] = get_arguments(rows[skill], tool, request)
    except (ValueError, json.JSONDecodeError, urllib.error.HTTPError) as e:
        record["seconds"] = time.perf_counter() - started
        record["result"] = f"model tidak menghasilkan panggilan tool yang sah ({e})"
        return record
    record["seconds"] = time.perf_counter() - started
    if not execute:
        record["ok"] = arguments_fit(case, record["arguments"])
        record["result"] = "tool tidak dijalankan (butuh jaringan atau file unggahan); yang dinilai hanya argumennya"
        return record
    namespace = tool.func.__globals__ if getattr(tool, "func", None) else {}
    for attr in ("DOWNLOADS_DIR", "CHARTS_DIR"):
        if attr in namespace:
            namespace[attr] = folder
    os.environ["PPTX_OUTPUT_DIR"] = str(folder)
    before = {p.name for p in folder.iterdir()}
    try:
        record["result"] = " ".join(str(tool.invoke(record["arguments"])).split())[:400]
    except Exception as e:  # noqa: BLE001 - the error is the result
        record["result"] = f"tool gagal: {type(e).__name__}: {str(e)[:200]}"
        return record
    record["files"] = new_files(folder, before)
    record["ok"] = "70" in record["result"] if case == "calculator" else bool(record["files"])
    return record


def arguments_fit(case: str, call_args: dict) -> bool:
    """For the tools that are not run: do the arguments carry what the request asked for?"""
    if case == "web_search":
        return bool(str(call_args.get("query") or "").strip())
    if case == "image":
        return call_args.get("file_id") == "t1/alasan.png"
    tasks = call_args.get("tasks")
    return isinstance(tasks, list) and len(tasks) >= 3 and all(isinstance(t, dict) and t.get("topic") for t in tasks)


# ---------------------------------------------------------------------------
# Runner cases
# ---------------------------------------------------------------------------

def run_script(command: list[str], cwd: Path) -> tuple[int, str]:
    done = subprocess.run(command, cwd=cwd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    return done.returncode, (done.stdout + done.stderr).strip()


def run_runner_case(case: str, folder: Path, model: str | None, args) -> dict:
    skill, _, generator, topic, output, builder, spec_name = RUNNER_CASES[case]
    folder.mkdir(parents=True, exist_ok=True)
    material = folder / "material.md"
    material.write_text("# Survei kebiasaan belanja online mahasiswa\n\n" + MATERIAL, encoding="utf-8")
    spec = folder / ("spec.md" if case == "mind_map" else "spec.json")
    before = {p.name for p in folder.iterdir()}
    started = time.perf_counter()
    if model is None:                                            # a ready spec (Claude): only the build script runs
        shutil.copy(OUT / "claude" / spec_name, spec)
        command = [sys.executable, str(POOL / skill / builder), str(spec), "-o", str(folder / output)]
    elif case == "mind_map":
        command = [sys.executable, str(POOL / skill / generator), str(material), "--model", model, "-o", str(folder / output),
                   "--save-outline", str(spec), "--timeout", str(args.timeout)]
    else:
        # --think is left at the generator's default: with a schema it asks without reasoning first
        command = [sys.executable, str(POOL / skill / generator), topic, "--context", str(material), "--model", model,
                   "-o", str(folder / output), "--save-json", str(spec), "--num-ctx", str(args.num_ctx),
                   "--timeout", str(args.timeout)]             # a slow machine is not a failed skill
    code, text = run_script(command, folder)
    head, fixes = said(text)
    return {"case": case, "ok": code == 0 and (folder / output).exists(), "seconds": time.perf_counter() - started,
            "arguments": spec.read_text(encoding="utf-8", errors="replace")[:6000] if spec.exists() else None,
            "result": head, "fixes": fixes, "files": new_files(folder, before)}


def said(text: str) -> tuple[str, list[str]]:
    """The result line of a skill script and the repairs it listed."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    head = next((ln for ln in lines if ln.startswith(("OK", "ERROR"))), lines[-1] if lines else "")
    return re.sub(r"created \S+[\\/]", "created ", head)[:300], [ln for ln in lines if ln.startswith("- ")][:8]


def rebuild_runner_case(case: str, folder: Path, saved: dict) -> dict:
    """Build again from the answer a model already gave (after a skill was repaired); the model is not asked."""
    skill, _, generator, topic, output, builder, _ = RUNNER_CASES[case]
    spec = folder / ("spec.md" if case == "mind_map" else "spec.json")
    shutil.rmtree(folder / "slides", ignore_errors=True)
    if case == "mind_map":
        command = [sys.executable, str(POOL / skill / builder), str(spec), "-o", str(folder / output)]
    else:
        command = [sys.executable, str(POOL / skill / generator), topic, "--context", str(folder / "material.md"),
                   "--from-json", str(spec), "-o", str(folder / output)]
    code, text = run_script(command, folder)
    head, fixes = said(text)
    files = sorted(p.name for p in folder.iterdir() if p.is_file() and p.name != "material.md")
    return {**saved, "ok": code == 0 and (folder / output).exists(), "result": head, "fixes": fixes, "files": files}


# ---------------------------------------------------------------------------
# One model
# ---------------------------------------------------------------------------

def run_model(model: str, args, rows: dict, only: list[str]) -> None:
    name = slug(model)
    base = OUT / name
    args.model = model
    args.thinking = smc.needs_thinking(args)
    try:
        smc.chat(args, [{"role": "user", "content": "Say: ready"}], tools=[{"type": "function", "function": {
            "name": "noop", "description": "Does nothing.", "parameters": {"type": "object", "properties": {}}}}], max_tokens=20)
        args.calls = "native"
    except urllib.error.HTTPError:
        args.calls = "schema"
    print(f"{model}: reasoning {'on' if args.thinking else 'off'}, tool calls {args.calls}", flush=True)
    results = load_results(name)
    results["_model"] = {"name": model, "reasoning": args.thinking, "calls": args.calls}
    for case in only:
        folder = base / case
        shutil.rmtree(folder, ignore_errors=True)
        if case in TOOL_CASES:
            record = run_tool_case(case, rows, folder, lambda row, tool, request: model_arguments(args, row, tool, request))
        else:
            record = run_runner_case(case, folder, model, args)
        results[case] = record
        save_results(name, results)
        print(f"  {'ok  ' if record['ok'] else 'FAIL'} {case:<14} {record['seconds']:5.0f}s  {record['result'][:110]}", flush=True)


def run_claude(rows: dict, only: list[str]) -> None:
    ready = json.loads((OUT / "claude" / "tools.json").read_text(encoding="utf-8"))
    for call_args in ready.values():                             # kept as plain JSON in the file, sent as a string like any tool call
        for key in ("content", "data"):
            if isinstance(call_args.get(key), (dict, list)):
                call_args[key] = json.dumps(call_args[key], ensure_ascii=False)
    results = load_results("claude")
    results["_model"] = {"name": "claude", "reasoning": None, "calls": "written in this session"}
    for case in only:
        folder = OUT / "claude-run" / case
        shutil.rmtree(folder, ignore_errors=True)
        if case in TOOL_CASES:
            record = run_tool_case(case, rows, folder, lambda row, tool, request: ready[TOOL_CASES[case][1]])
        else:
            record = run_runner_case(case, folder, None, None)
        record["seconds"] = None                                 # written by hand: there is no model time to report
        results[case] = record
        save_results("claude", results)
        print(f"  {'ok  ' if record['ok'] else 'FAIL'} {case:<14}  {record['result'][:110]}", flush=True)


def replay(model: str, rows: dict, only: list[str]) -> None:
    """Run the tools again with the arguments the model already wrote (after a tool was repaired).

    The model is not asked again, so the time of the first run is kept.
    """
    name = slug(model)
    results = load_results(name)
    for case in only:
        saved = results.get(case)
        if not saved or saved.get("arguments") is None:
            continue
        if case in RUNNER_CASES:
            results[case] = record = rebuild_runner_case(case, OUT / name / case, saved)
            save_results(name, results)
            print(f"  {'ok  ' if record['ok'] else 'FAIL'} {case:<14} rebuilt {record['result'][:110]}", flush=True)
            continue
        record = run_tool_case(case, rows, OUT / name / case, lambda row, tool, request: saved["arguments"])
        record["seconds"] = saved["seconds"]
        if record["ok"] and not saved["ok"]:
            record["first_result"] = saved.get("first_result") or saved["result"]     # what the tool said before the repair
        elif saved.get("first_result"):
            record["first_result"] = saved["first_result"]
        results[case] = record
        save_results(name, results)
        print(f"  {'ok  ' if record['ok'] else 'FAIL'} {case:<14} replay  {record['result'][:110]}", flush=True)


def load_results(name: str) -> dict:
    path = OUT / f"{name}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def save_results(name: str, results: dict) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / f"{name}.json").write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")


# ---------------------------------------------------------------------------
# Previews
# ---------------------------------------------------------------------------

EXPORT_SLIDES = r"""
param([string]$Pptx, [string]$OutDir, [int]$Width = 640)
$app = New-Object -ComObject PowerPoint.Application
$pres = $app.Presentations.Open($Pptx, $true, $false, $false)
$height = [int]($Width * $pres.PageSetup.SlideHeight / $pres.PageSetup.SlideWidth)
$i = 0
foreach ($slide in $pres.Slides) { $i++; $slide.Export((Join-Path $OutDir ("slide{0:D2}.png" -f $i)), "PNG", $Width, $height) }
$pres.Close()
$app.Quit()
"""


def export_slides(pptx: Path) -> list[Path]:
    """Slides of a .pptx as PNG files next to it (needs PowerPoint); [] when that is not possible."""
    target = pptx.parent / "slides"
    if target.exists() and any(target.iterdir()):
        return sorted(target.iterdir())
    target.mkdir(exist_ok=True)
    script = OUT / "_export_slides.ps1"
    script.write_text(EXPORT_SLIDES, encoding="utf-8")
    try:
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), str(pptx), str(target), "960"],
                       capture_output=True, timeout=240)
    except (OSError, subprocess.TimeoutExpired):
        return []
    return sorted(target.iterdir())


def embed(path: Path, width: int = 900) -> str:
    """The image as a data URI, so the page is one file that can be sent or opened anywhere."""
    import base64
    import io

    if path.suffix.lower() == ".svg":
        return "data:image/svg+xml;base64," + base64.b64encode(path.read_bytes()).decode("ascii")
    from PIL import Image

    image = Image.open(path).convert("RGB")
    if image.width > width:
        image = image.resize((width, round(image.height * width / image.width)), Image.LANCZOS)
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=78, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def docx_preview(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc, parts = Document(str(path)), []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            para = Paragraph(child, doc)
            text, style = html.escape(para.text.strip()), (para.style.name or "").lower()
            if not text:
                continue
            if "title" in style:
                parts.append(f"<h4>{text}</h4>")
            elif "heading" in style:
                parts.append(f"<h5>{text}</h5>")
            elif "list" in style:
                parts.append(f"<p class='li'>{text}</p>")
            else:
                parts.append(f"<p>{text}</p>")
        elif child.tag.endswith("}tbl"):
            rows = ["<tr>" + "".join(f"<td>{html.escape(c.text)}</td>" for c in row.cells) + "</tr>" for row in Table(child, doc).rows]
            parts.append("<table>" + "".join(rows) + "</table>")
    return "<div class='doc'>" + "".join(parts) + "</div>"


def xlsx_preview(path: Path) -> str:
    from openpyxl import load_workbook

    parts = []
    for sheet in load_workbook(str(path), data_only=True).worksheets:
        rows = ["<tr>" + "".join(f"<td>{html.escape('' if v is None else str(v))}</td>" for v in row) + "</tr>"
                for row in sheet.iter_rows(max_row=min(sheet.max_row, 12), values_only=True)]
        parts.append(f"<h5>Sheet: {html.escape(sheet.title)}</h5><table>" + "".join(rows) + "</table>")
    return "<div class='doc'>" + "".join(parts) + "</div>"


def preview(folder: Path, record: dict) -> str:
    """HTML that shows what the skill built for one case."""
    files = [folder / f for f in record.get("files", [])]
    pick = lambda *exts: next((f for f in files if f.suffix.lower() in exts and f.exists()), None)  # noqa: E731
    try:
        if pick(".pptx"):
            slides = export_slides(pick(".pptx"))
            if slides:
                return ("<div class='slides'>" + "".join(f"<img loading='lazy' src='{embed(s, 800)}' alt='slide {i}'>"
                                                         for i, s in enumerate(slides, 1)) + "</div>")
            return "<p class='muted'>File .pptx dibuat (PowerPoint tidak tersedia untuk pratinjau).</p>"
        if pick(".svg"):
            return f"<img loading='lazy' class='wide' src='{embed(pick('.svg'))}' alt='mind map'>"
        if pick(".png"):
            return f"<img loading='lazy' class='wide' src='{embed(pick('.png'), 1000)}' alt='hasil'>"
        if pick(".docx"):
            return docx_preview(pick(".docx"))
        if pick(".xlsx"):
            return xlsx_preview(pick(".xlsx"))
    except Exception as e:  # noqa: BLE001 - a preview must never stop the report
        return f"<p class='muted'>Pratinjau gagal: {html.escape(str(e))[:120]}</p>"
    return ""


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

PAGE = """<!doctype html>
<html lang="id"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Perbandingan hasil skill</title>
<style>
:root { --ink:#0b1f4d; --text:#33415c; --muted:#5d6b86; --line:#dfe4ef; --page:#f6f7fb; --card:#fff; --ok:#1f8a4c; --bad:#c2332b; --accent:#3f37c9; }
* { box-sizing: border-box; }
body { margin:0; background:var(--page); color:var(--text); font:15px/1.55 "Segoe UI", system-ui, sans-serif; }
main { max-width:1500px; margin:0 auto; padding:28px 20px 60px; }
h1 { color:var(--ink); font-size:26px; margin:0 0 6px; }
h2 { color:var(--ink); font-size:19px; margin:44px 0 4px; }
h4 { margin:0 0 6px; color:var(--ink); font-size:15px; } h5 { margin:10px 0 4px; color:var(--ink); font-size:13px; }
p { margin:0 0 8px; } .muted { color:var(--muted); font-size:13px; }
.request { background:var(--card); border-left:3px solid var(--accent); padding:8px 12px; margin:8px 0 14px; font-size:13.5px; white-space:pre-wrap; }
.cols { display:grid; gap:12px; grid-template-columns:repeat(auto-fit, minmax(225px, 1fr)); align-items:start; }
.cell { background:var(--card); border:1px solid var(--line); border-radius:8px; padding:12px; min-width:0; }
.head { display:flex; justify-content:space-between; gap:8px; align-items:baseline; margin-bottom:8px; }
.model { font-weight:600; color:var(--ink); } .time { font-size:12.5px; color:var(--muted); white-space:nowrap; }
.state { font-size:12px; font-weight:600; } .state.ok { color:var(--ok); } .state.bad { color:var(--bad); } .state.fixed { color:#9a6200; }
img.wide { width:100%; height:auto; border:1px solid var(--line); border-radius:6px; background:#fff; }
.slides { display:grid; grid-template-columns:repeat(auto-fit, minmax(170px, 1fr)); gap:6px; } .slides img { width:100%; border:1px solid var(--line); border-radius:4px; }
.doc { font-size:13px; border:1px solid var(--line); border-radius:6px; padding:10px 12px; background:#fff; overflow-x:auto; }
.doc .li::before { content:"\\2022  "; } table { border-collapse:collapse; margin:4px 0 8px; font-size:12.5px; }
td, th { border:1px solid var(--line); padding:3px 8px; text-align:left; vertical-align:top; }
.result { font-size:12.5px; color:var(--muted); margin-top:8px; overflow-wrap:anywhere; }
.note { font-size:13px; margin-top:8px; padding:6px 10px; background:#fff8e6; border-radius:6px; }
.cell img { cursor:zoom-in; }
#zoom { position:fixed; inset:0; background:rgba(11,31,77,.88); display:none; align-items:center; justify-content:center; padding:20px; cursor:zoom-out; z-index:9; }
#zoom img { max-width:100%; max-height:100%; background:#fff; border-radius:6px; }
details { margin-top:8px; font-size:12.5px; } summary { cursor:pointer; color:var(--accent); }
pre { margin:6px 0 0; padding:8px; background:#f1f3f9; border-radius:6px; overflow:auto; max-height:320px; font:12px/1.45 Consolas, monospace; white-space:pre-wrap; overflow-wrap:anywhere; }
.summary td:first-child { font-weight:600; color:var(--ink); } .summary { background:var(--card); width:100%; }
.summary td, .summary th { padding:6px 10px; font-size:13.5px; } .summary th { background:#eef0f8; color:var(--ink); }
</style></head><body><main>
<h1>Perbandingan hasil skill: 4B, 8B, dan Claude</h1>
<p class="muted">Setiap model menerima materi dan permintaan yang sama. Model lokal berjalan lewat Ollama di laptop (RTX 3060 6 GB).
Waktu Claude tidak dicantumkan karena argumennya ditulis langsung di sesi ini, bukan lewat server model.</p>
__SUMMARY__
<h2>Materi yang dipakai semua model</h2>
<div class="request">__MATERIAL__</div>
__SECTIONS__
</main>
<div id="zoom"><img alt="gambar diperbesar"></div>
<script>
const zoom = document.getElementById('zoom');
document.addEventListener('click', e => {
  if (e.target.closest('#zoom')) { zoom.style.display = 'none'; return; }
  if (e.target.matches('.cell img')) { zoom.firstElementChild.src = e.target.src; zoom.style.display = 'flex'; }
});
document.addEventListener('keydown', e => { if (e.key === 'Escape') zoom.style.display = 'none'; });
</script>
</body></html>
"""


def build_report() -> Path:
    names = [p.stem for p in sorted(OUT.glob("*.json")) if p.stem != "notes"]
    names.sort(key=lambda n: (n == "claude", "8b" in n, n))
    data = {n: load_results(n) for n in names}
    label = lambda n: LABELS.get(n, data[n].get("_model", {}).get("name", n))  # noqa: E731
    run_dir = lambda n: "claude-run" if n == "claude" else n  # noqa: E731
    titles = {**{k: v[2] for k, v in TOOL_CASES.items()}, **{k: v[1] for k, v in RUNNER_CASES.items()}}
    notes_file = OUT / "notes.json"                              # {case: {model: what a reader should notice}}, written after review
    notes = json.loads(notes_file.read_text(encoding="utf-8")) if notes_file.exists() else {}

    def state(record):
        if not record:
            return "<span class='muted'>belum dijalankan</span>"
        mark = "<span class='state bad'>gagal</span>"
        if record["ok"]:                                         # a tool that first refused the arguments is not a plain success
            mark = ("<span class='state fixed'>berhasil setelah tool diperbaiki</span>" if record.get("first_result")
                    else "<span class='state ok'>berhasil</span>")
        return mark + (f" <span class='time'>{record['seconds']:.0f} detik</span>" if record.get("seconds") else "")

    summary = ["<table class='summary'><tr><th>Skill</th>" + "".join(f"<th>{html.escape(label(n))}</th>" for n in names) + "</tr>"]
    sections = []
    for case in ORDER:
        if not any(case in data[n] for n in names):
            continue
        summary.append(f"<tr><td><a href='#{case}'>{html.escape(titles[case])}</a></td>" +
                       "".join(f"<td>{state(data[n].get(case))}</td>" for n in names) + "</tr>")
        request = TOOL_CASES[case][3] if case in TOOL_CASES else (RUNNER_CASES[case][3] or "Buat mind map dari dokumen materi di atas.")
        request = request.replace(MATERIAL, "[materi di atas]")
        cells = []
        for n in names:
            record = data[n].get(case)
            if not record:
                cells.append(f"<div class='cell'><div class='head'><span class='model'>{html.escape(label(n))}</span></div>"
                             "<p class='muted'>Belum dijalankan.</p></div>")
                continue
            shown = record.get("arguments")
            shown = shown if isinstance(shown, str) else json.dumps(shown, ensure_ascii=False, indent=1)
            written = f"<pre>{html.escape(shown or '(tidak ada)')}</pre>"
            picture = preview(OUT / run_dir(n) / case, record)
            if picture:                                          # with a picture the raw text is one click away
                written = f"<details><summary>Yang ditulis model</summary>{written}</details>"
            fixes = "".join(f"<br>{html.escape(f)}" for f in record.get("fixes") or [])
            if record.get("first_result"):
                fixes += f"<br>Sebelum tool diperbaiki: {html.escape(record['first_result'][:160])}"
            note = notes.get(case, {}).get(n, "")
            cells.append(
                f"<div class='cell'><div class='head'><span class='model'>{html.escape(label(n))}</span><span>{state(record)}</span></div>"
                f"{picture}{'' if picture else written}"
                f"<div class='result'>{html.escape(FOLDERS.sub('', record.get('result', '')))}{fixes}</div>"
                + (f"<p class='note'><b>Catatan:</b> {html.escape(note)}</p>" if note else "")
                + (written if picture else "") + "</div>")
        sections.append(f"<h2 id='{case}'>{html.escape(titles[case])}</h2><div class='request'>{html.escape(request)}</div>"
                        f"<div class='cols'>{''.join(cells)}</div>")
    summary.append("</table>")
    if notes.get("_page"):                                       # how the run was made, e.g. what else was using the GPU
        summary.insert(0, f"<p class='note'>{html.escape(notes['_page'])}</p>")
    page = (PAGE.replace("__SUMMARY__", "".join(summary)).replace("__MATERIAL__", html.escape(MATERIAL))
            .replace("__SECTIONS__", "".join(sections)))
    target = OUT / "index.html"
    target.write_text(page, encoding="utf-8")
    return target


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0], formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("models", nargs="*", help="local models, for example gemma3:4b qwen3:8b")
    ap.add_argument("--claude", action="store_true", help="run the specs in evals/comparison/claude/")
    ap.add_argument("--report", action="store_true", help="only rebuild the page")
    ap.add_argument("--replay", action="store_true",
                    help="do not ask the models again: run the tools with the arguments saved from the last run")
    ap.add_argument("--only", default="", help="case ids, comma separated: " + ", ".join(ORDER))
    ap.add_argument("--base-url", default="http://localhost:11434")
    ap.add_argument("--num-ctx", type=int, default=8192)
    ap.add_argument("--max-tokens", type=int, default=1500)
    ap.add_argument("--timeout", type=int, default=1800)
    args = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
    except Exception:  # noqa: BLE001
        pass

    only = [c for c in (args.only.split(",") if args.only else ORDER) if c in ORDER]
    if args.models or args.claude:
        sys.path.insert(0, str(POOL.parent))
        smc.install_app_stubs()
        rows = smc.inventory(list(smc.SKILLS))
        for model in args.models:
            replay(model, rows, only) if args.replay else run_model(model, args, rows, only)
        if args.claude:
            run_claude(rows, only)
    print(f"report: {build_report()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
