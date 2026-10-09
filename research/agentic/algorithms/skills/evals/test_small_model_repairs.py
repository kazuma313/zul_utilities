"""Repairs for what small models really wrote in evals/compare_models.py (gemma3:4b, 3 October 2026).

    python -m pytest research/agentic/algorithms/skills/evals -q -p no:cacheprovider
"""

import importlib.util
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
POOL = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(POOL.parent))
import small_model_check as smc  # noqa: E402

smc.install_app_stubs()


def load(path: str, name: str):
    """One script of a skill, under its own name (the skills share module names such as spec_normalizer)."""
    folder = str((POOL / path).parent)
    sys.path.insert(0, folder)
    try:
        spec = importlib.util.spec_from_file_location(name, POOL / path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(folder)
        for shared in ("spec_normalizer", "build_poster", "build_deck", "create_pptx", "charts", "icons", "font_metrics"):
            sys.modules.pop(shared, None)


MARKDOWN = """Ringkasan Eksekutif Survei

Survei ini meneliti kebiasaan belanja online mahasiswa di Medan.

**Temuan Utama:**
*   **Harga:** 54% responden memilih karena harga.
*   **Live streaming:** naik dari 12% menjadi 29%.

| Alasan    | Persentase (%) |
|-----------|----------------|
| Harga     | 54             |
| Kemudahan | 31             |
"""


def test_docx_reads_a_document_written_as_markdown(tmp_path):
    from utilities.docx import docx_skill

    spec = docx_skill._spec_from_markdown(MARKDOWN)
    assert spec["title"] == "Ringkasan Eksekutif Survei"
    assert [s["type"] for s in spec["sections"]] == ["paragraph", "heading1", "bullet", "table"]
    assert spec["sections"][1]["text"] == "Temuan Utama"
    assert spec["sections"][2]["items"][0] == "Harga: 54% responden memilih karena harga."
    assert spec["sections"][3] == {"type": "table", "headers": ["Alasan", "Persentase (%)"],
                                   "rows": [["Harga", "54"], ["Kemudahan", "31"]]}

    docx_skill.DOWNLOADS_DIR = tmp_path
    assert "created successfully" in docx_skill.create_docx.invoke({"content": MARKDOWN, "filename": "ringkasan.docx"})
    assert "Invalid JSON" in docx_skill.create_docx.invoke({"content": '{"title": "x" "sections": []}', "filename": "x.docx"})


def test_docx_closes_the_bracket_a_model_left_open(tmp_path):
    from utilities.docx import docx_skill

    cut = '{"title":"Ringkasan","sections":[{"type":"table","headers":["Alasan","Persen"],"rows":[["Harga","54%"]]}]'
    assert docx_skill._loads_lenient(cut)["sections"][0]["rows"] == [["Harga", "54%"]]
    assert docx_skill._closed('{"a": "text with } and [ inside') == '{"a": "text with } and [ inside"}'

    docx_skill.DOWNLOADS_DIR = tmp_path
    assert "created successfully" in docx_skill.create_docx.invoke({"content": cut, "filename": "ringkasan.docx"})


def test_xlsx_reads_rows_written_with_braces_and_no_keys(tmp_path):
    from openpyxl import load_workbook

    from utilities.xlsx import xlsx_skill

    data = '{\n "Alasan": [\n  {"alasan", "persen"},\n  {"harga", "54"},\n  {"kemudahan", "31"}\n ]\n}'
    assert xlsx_skill._loads_lenient(data) == {"Alasan": [["alasan", "persen"], ["harga", "54"], ["kemudahan", "31"]]}

    xlsx_skill.DOWNLOADS_DIR = tmp_path
    assert "created successfully" in xlsx_skill.create_xlsx.invoke({"data": data, "filename": "alasan.xlsx"})
    sheet = load_workbook(next(tmp_path.glob("*.xlsx")))["Alasan"]
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [["alasan", "persen"], ["harga", 54], ["kemudahan", 31]]


def test_xlsx_keeps_the_sheet_after_a_workbook_closed_too_early():
    from utilities.xlsx import xlsx_skill

    data = '{"Alasan":[{"alasan":"harga","persen":54}]}"Tren":[{"tahun":"2024","persen":12},{"tahun":"2026","persen":29}]}"'
    assert list(xlsx_skill._loads_lenient(data)) == ["Alasan", "Tren"]
    assert xlsx_skill._loads_lenient('{"Sales":[{"Month":"Jan","Revenue":50000}],"Summary":[{"Total":50000}]}') == {
        "Sales": [{"Month": "Jan", "Revenue": 50000}], "Summary": [{"Total": 50000}]}


def test_poster_keeps_a_long_stat_value_whole():
    normalizer = load("resaerch_poster/scripts/spec_normalizer.py", "poster_normalizer")
    warnings = []
    stats = normalizer._stats([{"label": "Belanja per bulan", "value": "Rp 310.000"},
                               {"label": "Terlalu panjang", "value": "lebih dari tiga ratus ribu rupiah"}], 6, warnings.append)
    assert stats[0]["value"] == "Rp 310.000"                    # was cut to "Rp 310.00", a different amount
    assert stats[1]["value"] == "lebih dari" and len(warnings) == 1


def test_poster_donut_keeps_percentages_that_do_not_reach_100():
    normalizer = load("resaerch_poster/scripts/spec_normalizer.py", "poster_normalizer")
    warnings = []
    section = {"chart": {"type": "donut", "categories": ["Harga", "Kemudahan"], "series": [{"name": "Alasan", "values": [54, 31]}]}}
    chart = normalizer._chart(section, warnings.append)
    assert chart["categories"] == ["Harga", "Kemudahan", "Others"]
    assert chart["series"][0]["values"] == [54, 31, 15] and chart["rest"] is True and warnings
    whole = normalizer._chart({"chart": {"type": "donut", "categories": ["A", "B"], "series": [{"values": [60, 40]}]}}, warnings.append)
    assert whole["categories"] == ["A", "B"] and whole["rest"] is False


SOURCE = "350 dari 500 mahasiswa di 5 kampus. 72% belanja online, rata-rata Rp 310.000. Harga 54%, kemudahan 31%, pilihan produk 15%."


def test_a_chart_keeps_only_numbers_from_the_source():
    deck = load("pptx_research/scripts/generate_deck.py", "research_generate")
    numbers = deck.source_numbers(SOURCE)
    assert {350.0, 500.0, 5.0, 72.0, 310000.0} <= numbers
    assert deck.from_source([54, 31, 15], numbers)
    assert deck.from_source([72, 28], numbers)                  # 28 is what is left of 100
    assert not deck.from_source([5, 7, 3], numbers)             # the example numbers of the old prompt
    assert not deck.from_source([0, 1, 0], numbers)

    raw = json.dumps({"slides": [
        {"layout": "cover", "title": "Paylater", "presenter": "Dr. Andi Rahman"},
        {"layout": "stat_chart", "title": "Temuan", "chart": {"series": [{"values": [5, 7, 3]}]}},
        {"layout": "chart", "title": "Alasan", "chart": {"series": [{"values": [54, 31, 15]}]}},
        {"layout": "closing", "title": "Terima Kasih", "subtitle": "Kontak: fakultas@unimed.ac.id | Sampai jumpa."}]})
    fixed, removed = deck.drop_invented(raw, SOURCE)
    slides = json.loads(fixed)["slides"]
    assert [s["title"] for s in slides] == ["Paylater", "Alasan", "Terima Kasih"] and len(removed) == 3
    assert slides[0]["presenter"] == "" and slides[2]["subtitle"] == "Sampai jumpa."


def test_a_deck_ends_at_its_conclusion_and_drops_an_invented_chart():
    deck = load("pptx_claude/scripts/generate_deck.py", "claude_generate")
    slides = [{"layout": "title", "title": "A"}, {"layout": "bullets", "title": "B"},
              {"layout": "chart", "title": "C", "chart_data": [{"values": [10, 20]}]},
              {"layout": "chart", "title": "D", "chart_data": [{"values": [54, 31, 15]}]},
              {"layout": "conclusion_cta", "title": "E"}, {"layout": "title", "title": "Terima Kasih"},
              {"layout": "bullets", "title": "Pertanyaan"}]
    raw, cut = deck.drop_slides_after_conclusion(json.dumps({"slides": slides}))
    assert cut == 2 and json.loads(raw)["slides"][-1]["title"] == "E"
    raw, removed = deck.drop_invented_charts(raw, SOURCE)
    assert [s["title"] for s in json.loads(raw)["slides"]] == ["A", "B", "D", "E"] and len(removed) == 1


def test_poster_footer_loses_a_contact_nobody_gave():
    generate = load("resaerch_poster/scripts/generate_poster.py", "poster_generate")
    raw = json.dumps({"footer": "Belajar Finansial | survei@kampus.ac.id | kampus.ac.id/survei"})
    fixed, dropped = generate.drop_invented_contacts(raw, "Survei kebiasaan belanja online mahasiswa di Medan")
    assert dropped == ["survei@kampus.ac.id", "kampus.ac.id/survei"]
    assert json.loads(fixed)["footer"] == ["Belajar Finansial"]
    assert generate.drop_invented_contacts(raw, "kontak: survei@kampus.ac.id, kampus.ac.id/survei")[1] == []
