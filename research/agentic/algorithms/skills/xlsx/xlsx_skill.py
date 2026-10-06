"""XLSX skill — read uploaded spreadsheets, deep-analyze CSV/XLSX, and create new Excel files."""

import io
import json
import logging
import re
import uuid
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns
from langchain_core.runnables import RunnableConfig
from langchain_core.tools import tool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.agent.minio_connection import upload_file_to_minio

logger = logging.getLogger(__name__)


def _loads_lenient(text: str):
    """json.loads that forgives what a small model adds around the JSON it passes as a string.

    Seen with qwen3:4b and qwen3:8b calling these tools: a stray quote or a remark after the closing
    brace, a code fence around the JSON, a comma before a closing bracket.  The first complete JSON
    value is taken and the rest is ignored.
    """
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip())
    start = min((i for i in (text.find("{"), text.find("[")) if i >= 0), default=0)
    text = text[start:]
    # qwen3:4b closed the workbook after its first sheet: {"Alasan":[...]}"Tren":[...]}".  Read as it is,
    # the first complete value is one sheet and the second sheet is lost without a word.  A "}" followed
    # by a key never occurs in valid JSON, so the sheets are joined again.
    text = re.sub(r'([\]}])\s*\}\s*"(?=[^"\\]+"\s*:\s*[\[{])', r'\1,"', text)
    error = None
    no_trailing = re.sub(r",\s*([\]}])", r"\1", text)
    # gemma3:4b writes a row as {"harga", "54"}: braces around values with no key.  That is a list.
    scalar = r'(?:"[^"\\]*"|-?\d+(?:\.\d+)?|true|false|null)'
    rows_as_lists = re.sub(r"\{(\s*" + scalar + r"\s*(?:,\s*" + scalar + r"\s*)*)\}", r"[\1]", no_trailing)
    for candidate in (text, no_trailing, rows_as_lists, _closed(rows_as_lists)):
        try:
            return json.JSONDecoder().raw_decode(candidate)[0]
        except json.JSONDecodeError as e:
            error = error or e
    raise error


def _closed(text: str) -> str:
    """The text plus the brackets it opened and never closed (a model often drops the last one)."""
    stack, in_string, escaped = [], False, False
    for ch in text:
        if in_string:
            if escaped:
                escaped = False
            elif ch == "\\":
                escaped = True
            elif ch == '"':
                in_string = False
        elif ch == '"':
            in_string = True
        elif ch in "{[":
            stack.append("}" if ch == "{" else "]")
        elif ch in "}]" and stack:
            stack.pop()
    return text + ('"' if in_string else "") + "".join(reversed(stack))


def _cell(value):
    """A number written as text ("54", "12.5") goes into the sheet as a number, so Excel can sum it."""
    if isinstance(value, str) and re.fullmatch(r"-?\d+(?:\.\d+)?", value.strip()):
        return float(value) if "." in value else int(value)
    return value

UPLOADS_DIR   = Path(__file__).parent.parent.parent / "app" / "uploads"
DOWNLOADS_DIR = Path(__file__).parent.parent.parent / "app" / "static" / "downloads"
CHARTS_DIR    = Path(__file__).parent.parent.parent / "app" / "static" / "charts"
DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)
CHARTS_DIR.mkdir(parents=True, exist_ok=True)

_EXCEL_EXTS = {".xlsx", ".xlsm", ".xls"}
_CSV_EXTS   = {".csv", ".tsv"}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _resolve_upload(file_id: str) -> Path:
    return UPLOADS_DIR / file_id


def _df_to_markdown(df: pd.DataFrame, max_rows: int = 50) -> str:
    """Return a markdown table; truncate long dataframes."""
    truncated = len(df) > max_rows
    sample = df.head(max_rows)
    lines = ["| " + " | ".join(str(c) for c in sample.columns) + " |",
             "| " + " | ".join(["---"] * len(sample.columns)) + " |"]
    for _, row in sample.iterrows():
        lines.append("| " + " | ".join(str(v) for v in row) + " |")
    result = "\n".join(lines)
    if truncated:
        result += f"\n\n*(showing first {max_rows} of {len(df)} rows)*"
    return result


def _style_header_row(ws, header_fill: str = "366092") -> None:
    """Bold white text on dark blue for the first row."""
    fill = PatternFill("solid", start_color=header_fill, end_color=header_fill)
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _auto_column_width(ws) -> None:
    """Set column widths based on content length."""
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)


# ── Tools ─────────────────────────────────────────────────────────────────────

@tool
def read_xlsx(file_id: str, sheet: str = "") -> str:
    """Read an uploaded Excel (.xlsx) or CSV file and return its contents.

    Use this when the user uploads a spreadsheet or CSV and asks to read,
    analyze, or summarize its data.

    Args:
        file_id: The file identifier from the upload (format: 'thread_id/uuid.xlsx').
        sheet:   Optional sheet name to read. Leave empty to read all sheets.
    """
    logger.info("[TOOL] read_xlsx called | file_id=%s | sheet=%r", file_id, sheet)

    path = _resolve_upload(file_id)
    if not path.exists():
        return f"File not found for file_id '{file_id}'. Make sure the file was uploaded."

    ext = path.suffix.lower()
    parts: list[str] = []

    try:
        if ext in _CSV_EXTS:
            sep = "\t" if ext == ".tsv" else ","
            df = pd.read_csv(str(path), sep=sep)
            parts.append(f"**{path.name}** — {len(df)} rows × {len(df.columns)} columns\n\n{_df_to_markdown(df)}")

        elif ext in _EXCEL_EXTS:
            all_sheets = pd.read_excel(str(path), sheet_name=None)
            target_sheets = {sheet: all_sheets[sheet]} if sheet and sheet in all_sheets else all_sheets

            if sheet and sheet not in all_sheets:
                avail = ", ".join(f'"{s}"' for s in all_sheets)
                return f"Sheet '{sheet}' not found. Available sheets: {avail}"

            for sname, df in target_sheets.items():
                parts.append(
                    f"### Sheet: {sname} — {len(df)} rows × {len(df.columns)} columns\n\n"
                    f"{_df_to_markdown(df)}"
                )
        else:
            return f"Unsupported file type '{ext}'. Supported: .xlsx, .xls, .xlsm, .csv, .tsv"

    except Exception as e:
        logger.error("[TOOL] read_xlsx error: %s", e)
        return f"Failed to read file: {e}"

    logger.info("[TOOL] read_xlsx done | sheets=%d", len(parts))
    return "\n\n---\n\n".join(parts)


@tool
def create_xlsx(data: str, filename: str, config: RunnableConfig = None) -> str:
    """Create a new Excel (.xlsx) file from JSON data and return a download URL.

    Use this when the user wants to export data to Excel, create a spreadsheet,
    or save tabular results as a downloadable file.

    Args:
        data: JSON string. Two formats accepted:
              - Array of objects (single sheet):
                '[{"Name":"Alice","Score":95},{"Name":"Bob","Score":82}]'
              - Object with sheet names as keys (multi-sheet):
                '{"Sales":[{"Month":"Jan","Revenue":50000}],"Summary":[{"Total":50000}]}'
        filename: Output filename (should end in .xlsx).
    """
    logger.info("[TOOL] create_xlsx called | filename=%s", filename)

    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    try:
        parsed = _loads_lenient(data)
    except Exception as e:
        return f"Invalid JSON data: {e}"

    # Normalise to dict of sheet_name → list[dict]
    if isinstance(parsed, list):
        sheets = {"Sheet1": parsed}
    elif isinstance(parsed, dict):
        sheets = parsed
    else:
        return "Data must be a JSON array or object."

    try:
        wb = Workbook()
        wb.remove(wb.active)  # remove default empty sheet

        for sheet_name, rows in sheets.items():
            ws = wb.create_sheet(title=str(sheet_name)[:31])

            if not rows:
                ws.append(["(empty)"])
                continue

            if isinstance(rows, list) and len(rows) > 1 and all(isinstance(r, list) for r in rows):
                # rows written as lists: the first one is the header
                df = pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]] if len(set(map(str, rows[0]))) == len(rows[0])
                                  and all(len(r) == len(rows[0]) for r in rows[1:]) else None)
            else:
                df = pd.DataFrame(rows)
            # Header row
            ws.append(list(df.columns))
            # Data rows
            for _, row in df.iterrows():
                ws.append([None if pd.isna(v) else _cell(v) for v in row])

            _style_header_row(ws)
            _auto_column_width(ws)
            # Freeze top row
            ws.freeze_panes = "A2"

        out_name = f"{uuid.uuid4().hex}_{filename}"
        out_path = DOWNLOADS_DIR / out_name
        wb.save(str(out_path))

        thread_id = (config or {}).get("configurable", {}).get("thread_id", "")
        username = thread_id.split("__")[0] if "__" in thread_id else "anonymous"
        minio_name = upload_file_to_minio(
            out_path.read_bytes(), username, filename,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            folder="xlsx",
        )
        if minio_name:
            url = f"/api/files/{minio_name}"
            suffix = ""
        else:
            url = f"/static/downloads/{out_name}"
            suffix = " (WARNING: MinIO upload failed — file may disappear on server restart)"
        logger.info("[TOOL] create_xlsx saved | url=%s", url)
        return f"Excel file created successfully. file_url:{url}{suffix}"

    except Exception as e:
        logger.error("[TOOL] create_xlsx error: %s", e)
        return f"Failed to create Excel file: {e}"


# ── CSV / XLSX deep analysis ───────────────────────────────────────────────────

def _save_chart(fig, label: str = "chart") -> str:
    """Save a matplotlib figure to CHARTS_DIR and return its chart_url.

    label is a short descriptive name (e.g. 'correlation_heatmap') prepended
    to the filename so the model knows what each chart contains.
    """
    fname = f"{label}_{uuid.uuid4().hex[:12]}.png"
    fig.savefig(str(CHARTS_DIR / fname), dpi=150, bbox_inches="tight")
    plt.close(fig)
    return f"chart_url:/static/charts/{fname}"


def _rotate_labels(ax) -> None:
    labels = [t.get_text() for t in ax.get_xticklabels()]
    if not labels:
        return
    avg_len = sum(len(l) for l in labels) / len(labels)
    effective = len(labels) * max(1.0, avg_len / 6)
    if effective > 12:
        ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=8)
    elif effective > 5:
        ax.set_xticklabels(labels, rotation=30, ha="right", fontsize=9)


def _analyze_dataframe(df: pd.DataFrame) -> tuple[str, list[str]]:
    """Run comprehensive analysis on a DataFrame. Returns (text_summary, chart_urls)."""
    sns.set_theme(style="darkgrid", palette="tab10")
    summary: list[str] = []
    charts: list[str] = []

    # ── Overview ──────────────────────────────────────────────────────────────
    summary.append(f"**Rows:** {df.shape[0]:,}  |  **Columns:** {df.shape[1]}")
    summary.append(f"**Columns:** {', '.join(df.columns.tolist())}")

    col_types = df.dtypes.value_counts()
    summary.append("**Column types:** " + ", ".join(f"{v}× {k}" for k, v in col_types.items()))

    # ── Missing data ──────────────────────────────────────────────────────────
    missing_total = df.isnull().sum().sum()
    if missing_total:
        missing_pct = missing_total / (df.shape[0] * df.shape[1]) * 100
        summary.append(f"\n**Missing values:** {missing_total:,} ({missing_pct:.1f}% of all cells)")
        col_missing = df.isnull().sum()
        col_missing = col_missing[col_missing > 0]
        for col, n in col_missing.items():
            summary.append(f"  - `{col}`: {n:,} ({n/len(df)*100:.1f}%)")
    else:
        summary.append("\n**Data quality:** ✓ No missing values")

    # ── Numeric analysis ──────────────────────────────────────────────────────
    numeric_cols = df.select_dtypes(include="number").columns.tolist()
    if numeric_cols:
        summary.append(f"\n**Numeric columns ({len(numeric_cols)}):** {', '.join(numeric_cols)}")
        desc = df[numeric_cols].describe().round(2)
        summary.append(desc.to_markdown())

        # Correlation heatmap
        if len(numeric_cols) > 1:
            corr = df[numeric_cols].corr()
            summary.append("\n**Top correlations:**")
            corr_pairs = (corr.where(~(corr == 1.0))
                          .stack().abs().sort_values(ascending=False).head(5))
            for (c1, c2), val in corr_pairs.items():
                summary.append(f"  - `{c1}` ↔ `{c2}`: {val:.2f}")

            fig, ax = plt.subplots(figsize=(max(6, len(numeric_cols)), max(5, len(numeric_cols) - 1)))
            sns.heatmap(corr, annot=True, fmt=".2f", cmap="coolwarm",
                        center=0, square=True, linewidths=0.5, ax=ax)
            ax.set_title("Correlation Heatmap", fontsize=13, fontweight="bold")
            plt.tight_layout()
            charts.append(_save_chart(fig, "correlation_heatmap"))
            logger.info("[TOOL] analyze_data correlation heatmap saved")

        # Distribution plots (up to 4 cols)
        plot_cols = numeric_cols[:4]
        n = len(plot_cols)
        fig, axes = plt.subplots(1, n, figsize=(5 * n, 4))
        if n == 1:
            axes = [axes]
        for ax, col in zip(axes, plot_cols):
            sns.histplot(df[col].dropna(), kde=True, ax=ax)
            ax.set_title(f"Distribution: {col}", fontsize=11)
            ax.set_xlabel(col)
        plt.tight_layout()
        charts.append(_save_chart(fig, "numeric_distributions"))
        logger.info("[TOOL] analyze_data distributions saved")

    # ── Categorical analysis ──────────────────────────────────────────────────
    cat_cols = [c for c in df.select_dtypes(include="object").columns
                if "id" not in c.lower() and df[c].nunique() <= 50]
    if cat_cols:
        summary.append(f"\n**Categorical columns ({len(cat_cols)}):** {', '.join(cat_cols)}")
        for col in cat_cols[:5]:
            vc = df[col].value_counts()
            top = vc.head(5)
            summary.append(f"\n`{col}` ({vc.nunique()} unique):")
            for val, cnt in top.items():
                summary.append(f"  - {val}: {cnt:,} ({cnt/len(df)*100:.1f}%)")

        # Bar charts for up to 4 categorical cols
        plot_cats = cat_cols[:4]
        n = len(plot_cats)
        fig, axes = plt.subplots(1, n, figsize=(6 * n, 5))
        if n == 1:
            axes = [axes]
        for ax, col in zip(axes, plot_cats):
            vc = df[col].value_counts().head(10)
            sns.barplot(x=vc.values, y=vc.index.astype(str), ax=ax,
                        palette="tab10", errorbar=None)
            ax.set_title(f"Top values: {col}", fontsize=11)
            ax.set_xlabel("Count")
        plt.tight_layout()
        charts.append(_save_chart(fig, "categorical_bars"))
        logger.info("[TOOL] analyze_data categorical chart saved")

    # ── Time series analysis ──────────────────────────────────────────────────
    date_cols = [c for c in df.columns
                 if any(k in c.lower() for k in ("date", "time", "timestamp", "created", "updated"))]
    if date_cols:
        date_col = date_cols[0]
        df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
        df = df.dropna(subset=[date_col])
        if not df.empty:
            span = (df[date_col].max() - df[date_col].min()).days
            summary.append(f"\n**Time range ({date_col}):** "
                           f"{df[date_col].min().date()} → {df[date_col].max().date()} "
                           f"({span} days)")

            if numeric_cols:
                plot_num = numeric_cols[:3]
                fig, axes = plt.subplots(len(plot_num), 1,
                                         figsize=(12, 4 * len(plot_num)))
                if len(plot_num) == 1:
                    axes = [axes]
                for ax, col in zip(axes, plot_num):
                    grouped = df.groupby(date_col)[col].mean()
                    ax.plot(grouped.index, grouped.values, linewidth=2, marker="o", markersize=3)
                    ax.set_title(f"{col} over time (daily avg)", fontsize=11)
                    ax.set_xlabel(date_col)
                    ax.set_ylabel(col)
                    ax.grid(True, alpha=0.3)
                    _rotate_labels(ax)
                plt.tight_layout()
                charts.append(_save_chart(fig, "timeseries_trends"))
                logger.info("[TOOL] analyze_data time-series chart saved")

    return "\n".join(summary), charts


@tool
def analyze_data(file_id: str) -> str:
    """Run a comprehensive automatic analysis on an uploaded CSV or Excel file.

    Immediately generates:
    - Data overview (rows, columns, types, missing values)
    - Statistical summary of all numeric columns
    - Top correlations + correlation heatmap chart
    - Distribution histograms for numeric columns
    - Bar charts for categorical columns
    - Time-series trend charts (if date columns are detected)

    Use this whenever a user uploads a CSV/Excel file and wants insights,
    a summary, or visualizations — DO NOT ask what they want, just analyze it.

    Args:
        file_id: The file identifier from the upload (format: 'thread_id/uuid.csv').
    """
    logger.info("[TOOL] analyze_data called | file_id=%s", file_id)

    path = UPLOADS_DIR / file_id
    if not path.exists():
        return f"File not found for file_id '{file_id}'."

    ext = path.suffix.lower()
    try:
        if ext in {".csv", ".tsv"}:
            sep = "\t" if ext == ".tsv" else ","
            df = pd.read_csv(str(path), sep=sep)
        elif ext in {".xlsx", ".xls", ".xlsm"}:
            df = pd.read_excel(str(path))
        else:
            return f"Unsupported file type '{ext}' for analysis. Use CSV or XLSX."
    except Exception as e:
        return f"Failed to load file: {e}"

    logger.info("[TOOL] analyze_data loaded | shape=%s", df.shape)

    text, chart_urls = _analyze_dataframe(df)

    result_parts = [f"## Analysis of `{path.name}`\n", text]

    if chart_urls:
        # Build a labelled list so the model knows what each chart shows
        labelled = []
        for url in chart_urls:
            fname = url.split("/")[-1]          # e.g. correlation_heatmap_abc123.png
            label = fname.rsplit("_", 1)[0].replace("_", " ").title()  # "Correlation Heatmap"
            labelled.append(f"- {label}: {url}")

        chart_section = (
            "\n\n---\n"
            "## GENERATED CHARTS — USE THESE URLS IN REPORTS\n"
            "Copy the chart_url values EXACTLY into image_url fields when building PPT or Word reports.\n"
            "Each chart name tells you what it shows:\n\n"
            + "\n".join(labelled)
            + "\n\n"
            "For PPT: use layout='image' or 'image_bullets' and set image_url to the chart_url above.\n"
            "For Word: use type='image' and set image_url to the chart_url above.\n"
            "---"
        )
        result_parts.append(chart_section)

    logger.info("[TOOL] analyze_data done | charts=%d", len(chart_urls))
    return "\n\n".join(result_parts)
