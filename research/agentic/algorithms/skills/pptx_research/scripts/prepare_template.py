"""Make `assets/template.pptx` from your original template file.

The original export is about 25 MB because a few photos are 5000+ pixels wide.
This script writes a copy whose photos are at most 2000 px (same look, about
a tenth of the size), so every generated deck stays small.  The original file
is not changed.

    python scripts/prepare_template.py "C:/path/Black and White Modern Research Proposal Presentation.pptx"
"""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("source", help="The original template .pptx")
    ap.add_argument("-o", "--output", default=str(Path(__file__).resolve().parent.parent / "assets" / "template.pptx"))
    ap.add_argument("--max-px", type=int, default=2000)
    args = ap.parse_args()
    try:
        from PIL import Image
        from pptx import Presentation
    except ImportError:
        print("ERROR: run `pip install python-pptx` first")
        return 1
    src = Path(args.source)
    if not src.is_file():
        print(f"ERROR: file not found: {src}")
        return 1
    prs = Presentation(str(src))
    if len(prs.slides) != 15:
        print(f"ERROR: expected the 15-slide research proposal template, found {len(prs.slides)} slides")
        return 1
    shrunk = 0
    for slide in prs.slides:
        for blip in slide.element.iter(A_BLIP):
            rid = blip.get(R_EMBED)
            if not rid:
                continue
            part = slide.part.related_part(rid)
            if len(part.blob) < 400_000:
                continue
            im = Image.open(io.BytesIO(part.blob))
            im.thumbnail((args.max_px, args.max_px))
            buf = io.BytesIO()
            if im.mode in ("RGBA", "LA", "P") and "transparency" in im.info or im.mode in ("RGBA", "LA"):
                alpha = im.convert("RGBA").getchannel("A")
                if alpha.getextrema()[0] < 255:          # really transparent: keep PNG
                    im.convert("RGBA").save(buf, "PNG", optimize=True)
                else:
                    im.convert("RGB").save(buf, "JPEG", quality=85, optimize=True)
            else:
                im.convert("RGB").save(buf, "JPEG", quality=85, optimize=True)
            buf.seek(0)
            _, new_rid = slide.part.get_or_add_image_part(buf)
            blip.set(R_EMBED, new_rid)
            if not any(b.get(R_EMBED) == rid for b in slide.element.iter(A_BLIP)):
                slide.part.drop_rel(rid)      # otherwise the big original stays in the file
            shrunk += 1
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"OK: wrote {out} ({out.stat().st_size / 1e6:.1f} MB, {shrunk} photos resized; "
          f"original was {src.stat().st_size / 1e6:.1f} MB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
