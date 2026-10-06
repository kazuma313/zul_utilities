"""pptx-research engine: fill the research-proposal template from a JSON spec.

The design is NOT drawn by this code.  It lives in `assets/template.pptx`
(your copy of the template).  This script copies the template slides you ask
for, puts them in the standard section order, and replaces text, pictures
and charts.  Requires: Python 3.9+ and `python-pptx` (which brings Pillow).

    python scripts/build_deck.py deck.json -o proposal.pptx
    python scripts/build_deck.py deck.json -o proposal.pptx --keep-order --no-agenda --json

Python:
    from build_deck import build_deck, format_result
    result = build_deck(spec, "proposal.pptx")      # never raises
"""

from __future__ import annotations

import argparse
import copy
import io
import json
import math
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from font_metrics import text_width_in  # noqa: E402
from spec_normalizer import normalize_spec, parse_slides  # noqa: E402

SCRIPTS_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPTS_DIR.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "template.pptx"
TEMPLATE_ENV = "PPTX_RESEARCH_TEMPLATE"
OUTPUT_DIR_ENV = "PPTX_OUTPUT_DIR"
IMAGE_DIRS_ENV = "PPTX_IMAGE_DIRS"

EMU = 914400
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
R_NS = "{%s}" % NS["r"]

# Template slide number for each layout
TEMPLATE_SLIDE = {
    "cover": 1, "intro": 2, "agenda": 3, "section": 4, "three_cards": 5, "team": 6,
    "two_cards": 7, "people": 8, "chart": 9, "timeline": 10, "analysis": 11,
    "stat_chart": 12, "gallery": 13, "testimonials": 14, "closing": 15,
}

# Bold display font: average glyph width as a fraction of the font size
TITLE_GLYPH = 0.56


class Ctx:
    """Per-build state: warnings and image search folders."""

    def __init__(self, image_dirs):
        self.warnings: list[str] = []
        self.image_dirs = [Path(d) for d in (image_dirs or [])]
        for d in os.environ.get(IMAGE_DIRS_ENV, "").split(os.pathsep):
            if d.strip():
                self.image_dirs.append(Path(d.strip()))
        self.image_dirs.append(Path.cwd())
        self.tag = ""

    def warn(self, msg: str):
        text = f"{self.tag}: {msg}" if self.tag else msg
        if text not in self.warnings:
            self.warnings.append(text)


# ---------------------------------------------------------------------------
# Slide copy / delete
# ---------------------------------------------------------------------------

def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            return layout
    return prs.slide_layouts[len(prs.slide_layouts) - 1]


def duplicate_slide(prs, src):
    """Append a copy of `src` (shapes, background, pictures) and return it."""
    dst = prs.slides.add_slide(_blank_layout(prs))
    # Fill the existing <p:spTree> in place: python-pptx caches `slide.shapes`
    # on it, so swapping the whole element would leave that cache pointing nowhere.
    new_csld = dst.element.cSld
    tree = dst.shapes._spTree
    for child in list(tree):
        tree.remove(child)
    for child in src.shapes._spTree:
        tree.append(copy.deepcopy(child))
    for bg in new_csld.findall("p:bg", NS):
        new_csld.remove(bg)
    src_bg = src.element.cSld.find("p:bg", NS)
    if src_bg is not None:
        new_csld.insert(0, copy.deepcopy(src_bg))

    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            rid_map[rid] = dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = dst.part.relate_to(rel.target_part, rel.reltype)
    if rid_map:
        for el in new_csld.iter():
            for attr, val in list(el.attrib.items()):
                if attr.startswith(R_NS) and val in rid_map:
                    el.set(attr, "__new__" + rid_map[val])   # two passes: avoid rId1 -> rId2 -> rId3 chains
        for el in new_csld.iter():
            for attr, val in list(el.attrib.items()):
                if val.startswith("__new__"):
                    el.set(attr, val[7:])
    return dst


def delete_slides(prs, slides):
    id_list = prs.slides._sldIdLst
    targets = {s.slide_id for s in slides}
    for sld_id in list(id_list):
        if sld_id.id in targets:
            prs.part.drop_rel(sld_id.rId)
            id_list.remove(sld_id)


# ---------------------------------------------------------------------------
# Shape helpers
# ---------------------------------------------------------------------------

def shape(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"template shape {name!r} not found - was assets/template.pptx edited?")


def remove(slide, *names):
    for name in names:
        try:
            sh = shape(slide, name)
        except KeyError:
            continue
        rids = sh._element.xpath(".//a:blip/@r:embed")
        sh._element.getparent().remove(sh._element)
        for rid in rids:
            _drop_unused_rel(slide, rid)


def _inches(sh):
    return sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU


def _clean(text) -> str:
    text = "" if text is None else str(text)
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    return text.strip()


def _scale_fonts(tx_body, scale: float):
    if abs(scale - 1.0) < 0.001:
        return
    for rpr in tx_body.iter("{%s}rPr" % NS["a"], "{%s}endParaRPr" % NS["a"]):
        if rpr.get("sz"):
            rpr.set("sz", str(max(800, int(int(rpr.get("sz")) * scale))))
        if rpr.get("spc"):
            rpr.set("spc", str(int(int(rpr.get("spc")) * scale)))
    for spc in tx_body.iter("{%s}spcPts" % NS["a"]):
        spc.set("val", str(max(800, int(int(spc.get("val")) * scale))))


def _write(sh, lines: list[str]):
    """Replace the text of a text box, keeping the first run's formatting.

    One paragraph per entry of `lines` (each a copy of the template's first
    paragraph, so bullets / alignment / spacing survive)."""
    tx = sh._element.find(".//p:txBody", NS)
    if tx is None:
        tx = sh._element.find(".//{%s}txBody" % NS["p"])
    paras = tx.findall("a:p", NS)
    first = paras[0]
    runs = first.findall("a:r", NS)
    if not runs:  # empty template paragraph: build a run from endParaRPr
        run = first.makeelement("{%s}r" % NS["a"], {})
        t = first.makeelement("{%s}t" % NS["a"], {})
        run.append(t)
        first.append(run)
        runs = [run]
    for extra in runs[1:]:
        first.remove(extra)
    for br in first.findall("a:br", NS):
        first.remove(br)
    for p in paras[1:]:
        tx.remove(p)
    proto = copy.deepcopy(first)
    lines = lines or [""]
    first.find("a:r", NS).find("a:t", NS).text = lines[0]
    prev = first
    for line in lines[1:]:
        p = copy.deepcopy(proto)
        p.find("a:r", NS).find("a:t", NS).text = line
        prev.addnext(p)
        prev = p
    return tx


def _font_pt(sh) -> float:
    sz = sh._element.xpath(".//a:rPr/@sz")
    return int(sz[0]) / 100 if sz else 18.0


def _line_pt(sh) -> float:
    ln = sh._element.xpath(".//a:lnSpc/a:spcPts/@val")
    return int(ln[0]) / 100 if ln else _font_pt(sh) * 1.2


def _face(sh) -> str:
    face = (sh._element.xpath(".//a:latin/@typeface") or ["HK Grotesk"])[0].lower()
    if "sauce" in face:
        return "display" if "bold" in face else "display_light"
    return "body_bold" if "bold" in face else "body"


def _spacing_pt(sh) -> float:
    spc = sh._element.xpath(".//a:rPr/@spc")
    return int(spc[0]) / 100 if spc else 0.0


def _wrap_lines(text: str, width_in: float, size: float, face: str, spc: float) -> float:
    """Greedy word wrap with real glyph widths; inf when one word cannot fit."""
    lines, cur = 1, ""
    for word in text.split():
        if text_width_in(word, size, face, spc) > width_in:
            return math.inf
        trial = f"{cur} {word}".strip()
        if text_width_in(trial, size, face, spc) <= width_in:
            cur = trial
        else:
            lines, cur = lines + 1, word
    return lines


def set_title(ctx: Ctx, sh, text, avail_h: float, floor: float = 0.4, grow_up: bool = False, **_):
    """Big display text: the largest size whose wrapped height fits `avail_h` inches."""
    text = re.sub(r"\s+", " ", _clean(text)) or " "
    _, top, w_in, h_in = _inches(sh)
    size, line, face, spc = _font_pt(sh), _line_pt(sh), _face(sh), _spacing_pt(sh)
    usable = w_in * 0.98
    scale, lines, s = None, 1, 1.0
    while s >= floor - 1e-9:
        n = _wrap_lines(text, usable, size * s, face, spc * s)
        if n * line * s / 72 <= avail_h + 1e-6:
            scale, lines = s, n
            break
        s -= 0.05
    if scale is None:
        scale = floor
        n = _wrap_lines(text, usable, size * floor, face, spc * floor)
        lines = 3 if n == math.inf else n
        ctx.warn(f"title {text[:40]!r} is too long for this slide - it may overflow (shorten it)")
    tx = _write(sh, [text])
    _scale_fonts(tx, scale)
    new_h = lines * line * scale / 72
    if grow_up:  # keep the bottom edge where the designer put it
        sh.top = int((top + h_in - new_h) * EMU)
    sh.height = int(new_h * EMU)
    return scale


def set_text(ctx: Ctx, sh, text, capacity: int | None = None, floor: float = 0.7, label: str = "text",
             grow: float = 1.3):
    """Body text.  `capacity` = characters that fit at full size (defaults to the
    length of the template's own placeholder text, which the designer sized the box for)."""
    lines = [_clean(x) for x in (text if isinstance(text, (list, tuple)) else [text])]
    lines = [x for x in lines if x] or [""]
    if len(lines) == 1:
        lines = [x.strip() for x in lines[0].split("\n") if x.strip()] or [""]
    if capacity is None:
        capacity = max(12, len(sh.text_frame.text))
    # each extra paragraph costs about half a line
    length = sum(len(x) for x in lines) + 40 * (len(lines) - 1)
    scale = 1.0
    if length < capacity * 0.55 and grow:
        scale = min(grow, math.sqrt(capacity * 0.8 / max(length, 1)))
    if length > capacity:
        scale = max(floor, math.sqrt(capacity / length))
        if length > capacity / (floor * floor) * 1.05:
            ctx.warn(f"{label} has {length} characters, about {int(capacity / (floor * floor))} fit - "
                     "it may overflow (shorten it)")
    tx = _write(sh, lines)
    _scale_fonts(tx, scale)
    return scale


def set_line(ctx: Ctx, sh, text, floor: float = 0.6, label: str = "label", **_):
    """Single-line label: shrink so it stays on one line of its box."""
    text = re.sub(r"\s+", " ", _clean(text))
    _, _, w_in, _ = _inches(sh)
    need = text_width_in(text, _font_pt(sh), _face(sh), _spacing_pt(sh)) * 1.03
    scale = 1.0 if need <= w_in else max(floor, w_in / need)
    if need * floor > w_in * 1.05:
        ctx.warn(f"{label} {text[:40]!r} is too long for its box (shorten it)")
    tx = _write(sh, [text or " "])
    _scale_fonts(tx, scale)


# ---------------------------------------------------------------------------
# Pictures
# ---------------------------------------------------------------------------

def _find_image(ctx: Ctx, ref: str) -> Path | None:
    if not ref:
        return None
    ref = ref.strip()
    if re.match(r"^https?://", ref, re.I):
        ctx.warn(f"image {ref!r} is a web address - only local files are used")
        return None
    if ref.startswith("file://"):
        ref = ref[7:]
    p = Path(ref).expanduser()
    if p.is_absolute():
        return p if p.is_file() else None
    for d in ctx.image_dirs:
        for cand in (d / ref, d / Path(ref).name):
            if cand.is_file():
                return cand
    return None


def _prepared_image(path: Path, max_px: int = 2000):
    """(file-like JPEG/PNG, aspect).  Large photos are downscaled to keep decks small."""
    from PIL import Image, ImageOps
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        if max(im.size) > max_px:
            im.thumbnail((max_px, max_px))
        has_alpha = im.mode in ("RGBA", "LA", "P")
        buf = io.BytesIO()
        if has_alpha:
            im.convert("RGBA").save(buf, "PNG", optimize=True)
        else:
            im.convert("RGB").save(buf, "JPEG", quality=88, optimize=True)
        buf.seek(0)
        return buf, im.size[0] / im.size[1]


def _avatar():
    """Neutral silhouette used when a person has no photo (never a stranger's face)."""
    from PIL import Image, ImageDraw
    size = 600
    im = Image.new("RGB", (size, size), "#2B2B2B")
    d = ImageDraw.Draw(im)
    for y in range(size):  # soft vertical gradient like the template cards
        c = int(28 + 40 * y / size)
        d.line([(0, y), (size, y)], fill=(c, c, c))
    d.ellipse([205, 120, 395, 310], fill="#8C8C8C")
    d.ellipse([90, 350, 510, 900], fill="#8C8C8C")
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    buf.seek(0)
    return buf, 1.0


def _picture_shape(group):
    """The freeform with a picture fill inside a template group (+ its rendered aspect)."""
    el = group._element
    sp = el.xpath(".//p:sp[p:spPr/a:blipFill]")
    if not sp:
        return None, 1.0
    sp = sp[0]
    ext = sp.find("p:spPr/a:xfrm/a:ext", NS)
    gx = el.find("p:grpSpPr/a:xfrm", NS)
    sx = sy = 1.0
    if gx is not None and gx.find("a:chExt", NS) is not None:
        ch, ge = gx.find("a:chExt", NS), gx.find("a:ext", NS)
        if int(ch.get("cx")) and int(ch.get("cy")):
            sx, sy = int(ge.get("cx")) / int(ch.get("cx")), int(ge.get("cy")) / int(ch.get("cy"))
    w, h = int(ext.get("cx")) * sx, int(ext.get("cy")) * sy
    return sp, (w / h if h else 1.0)


def _drop_unused_rel(slide, rid):
    """Drop a picture relationship nobody uses any more, so the old photo is not saved."""
    if not rid:
        return
    for el in slide.element.iter():
        for attr, val in el.attrib.items():
            if attr.startswith(R_NS) and val == rid:
                return
    try:
        slide.part.drop_rel(rid)
    except KeyError:
        pass


def set_picture(ctx: Ctx, slide, group_name: str, ref: str | None, person: bool = False, label: str = "image"):
    """Swap the photo inside a template group, cropping to cover (no stretching).

    No usable file: people get a neutral avatar, other slots keep the template photo."""
    try:
        group = shape(slide, group_name)
    except KeyError:
        return
    sp, slot_aspect = _picture_shape(group)
    if sp is None:
        return
    path = _find_image(ctx, ref) if ref else None
    if ref and not path:
        ctx.warn(f"{label} {ref!r} not found - " + ("placeholder avatar used" if person else "template photo kept"))
    if path:
        try:
            stream, aspect = _prepared_image(path)
        except Exception as e:
            ctx.warn(f"{label} {ref!r} could not be read ({e})")
            path = None
    if not path:
        if not person:
            return
        stream, aspect = _avatar()
    _, rid = slide.part.get_or_add_image_part(stream)
    blip_fill = sp.find("p:spPr/a:blipFill", NS)
    blip = blip_fill.find("a:blip", NS)
    old_rid = blip.get(R_NS + "embed")
    blip.set(R_NS + "embed", rid)
    _drop_unused_rel(slide, old_rid)
    stretch = blip_fill.find("a:stretch", NS)
    rect = stretch.find("a:fillRect", NS) if stretch is not None else None
    if rect is not None:
        for k in ("l", "t", "r", "b"):
            rect.attrib.pop(k, None)
        if aspect > slot_aspect:      # picture is wider than the slot: crop left/right
            over = int((aspect / slot_aspect - 1) / 2 * 100000)
            rect.set("l", str(-over)); rect.set("r", str(-over))
        elif aspect < slot_aspect:    # taller: crop top/bottom
            over = int((slot_aspect / aspect - 1) / 2 * 100000)
            rect.set("t", str(-over)); rect.set("b", str(-over))


# ---------------------------------------------------------------------------
# Charts (native, editable) in the template's black / white / grey palette
# ---------------------------------------------------------------------------

def add_chart(ctx: Ctx, slide, picture_name: str, chart: dict | None, default_type: str, dark: bool,
              box: tuple | None = None):
    """Replace the template's chart *picture* by a real chart.  No data -> picture removed."""
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Pt

    try:
        pic = shape(slide, picture_name)
    except KeyError:
        pic = None
    if pic is not None:
        x, y, w, h = pic.left, pic.top, pic.width, pic.height
        rids = pic._element.xpath(".//a:blip/@r:embed")
        pic._element.getparent().remove(pic._element)
        for rid in rids:
            _drop_unused_rel(slide, rid)
    if not chart:
        return None
    if box:
        x, y, w, h = (int(v * EMU) for v in box)

    kind = chart.get("type") or default_type
    types = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "radar": XL_CHART_TYPE.RADAR_FILLED, "area": XL_CHART_TYPE.AREA,
    }
    data = CategoryChartData()
    data.categories = chart["categories"]
    series = chart["series"][:1] if kind in ("pie", "doughnut") else chart["series"]
    for ser in series:
        data.add_series(ser["name"], ser["values"])
    frame = slide.shapes.add_chart(types.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED), Emu(x), Emu(y), Emu(w), Emu(h), data)
    ch = frame.chart

    ink = RGBColor(0xFF, 0xFF, 0xFF) if dark else RGBColor(0x00, 0x00, 0x00)
    grid = RGBColor(0x3A, 0x3A, 0x3A) if dark else RGBColor(0xD9, 0xD9, 0xD9)
    palette = (["5B5B5B", "FFFFFF", "A6A6A6", "D9D9D9", "7F7F7F", "404040"] if dark
               else ["5B5B5B", "000000", "A6A6A6", "D9D9D9", "7F7F7F", "404040"])

    ch.font.size = Pt(18)
    ch.font.name = "HK Grotesk"
    ch.font.color.rgb = ink
    ch.has_title = False
    ch.has_legend = len(series) > 1 or kind in ("pie", "doughnut")
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(18)
        ch.legend.font.color.rgb = ink

    if kind in ("pie", "doughnut"):
        plot = ch.plots[0]
        plot.has_data_labels = True
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
        plot.data_labels.font.size = Pt(16)
        plot.data_labels.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F) if not dark else ink
        for i, point_value in enumerate(series[0]["values"]):
            pt = plot.series[0].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(palette[i % len(palette)])
    else:
        for i, ser in enumerate(ch.plots[0].series):
            color = RGBColor.from_string(palette[i % len(palette)])
            if kind == "line":
                ser.format.line.color.rgb = color
                ser.format.line.width = Pt(3)
                ser.smooth = True
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = color
                ser.marker.format.line.color.rgb = color
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = color
                if kind in ("radar", "area"):
                    ser.format.line.color.rgb = color
                    clr = ser.format.fill._xPr.find(".//a:solidFill/a:srgbClr", NS)
                    if clr is not None:           # see-through so overlapping series stay visible
                        alpha = clr.makeelement("{%s}alpha" % NS["a"], {"val": "60000"})
                        clr.append(alpha)
        if kind in ("bar", "column"):
            ch.plots[0].gap_width = 60
            ch.plots[0].overlap = -10
        if kind != "radar":
            for axis in (ch.category_axis, ch.value_axis):
                axis.format.line.color.rgb = grid
                axis.tick_labels.font.size = Pt(16)
                axis.tick_labels.font.color.rgb = ink
            ch.value_axis.has_major_gridlines = True
            ch.value_axis.major_gridlines.format.line.color.rgb = grid
            ch.category_axis.has_major_gridlines = False
    return frame


# ---------------------------------------------------------------------------
# Layout fillers.  `s` is a normalised slide dict.
# ---------------------------------------------------------------------------

def _opt_text(ctx, slide, name, value, **kw):
    """Fill a body text box, or remove it when the spec gives nothing."""
    if value:
        set_text(ctx, shape(slide, name), value, **kw)
    else:
        remove(slide, name)


def fill_cover(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Research Proposal", avail_h=1.98, floor=0.45,
              grow_up=True, glyph=0.54)
    if s.get("subtitle"):
        shape(sl, "TextBox 6").width = int((7.1 if s.get("tagline") else 16.2) * EMU)  # left of the pill
        set_line(ctx, shape(sl, "TextBox 6"), s["subtitle"], floor=0.45, label="subtitle")
    else:
        remove(sl, "TextBox 6")
    if s.get("tagline"):
        set_line(ctx, shape(sl, "TextBox 10"), s["tagline"], label="tagline")
    else:
        remove(sl, "TextBox 10", "Group 7")
    if s.get("organization"):
        set_line(ctx, shape(sl, "TextBox 11"), s["organization"], label="organization")
    else:
        remove(sl, "TextBox 11")
    if s.get("presenter"):
        set_line(ctx, shape(sl, "TextBox 12"), s.get("presenter_label") or "Presented By:")
        set_line(ctx, shape(sl, "TextBox 13"), s["presenter"], label="presenter")
    else:
        remove(sl, "TextBox 12", "TextBox 13")


def fill_intro(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Hello !", avail_h=2.1)
    title_w = shape(sl, "TextBox 5").width
    shape(sl, "TextBox 5").width = int(5.0 * EMU) if len(s.get("title") or "Hello !") <= 8 else title_w
    _opt_text(ctx, sl, "TextBox 10", s.get("lead"), label="lead")
    paras = s.get("paragraphs") or []
    if len(paras) == 1:
        set_text(ctx, shape(sl, "TextBox 9"), paras[0], capacity=1150, label="paragraphs")
        remove(sl, "TextBox 11")
    else:
        _opt_text(ctx, sl, "TextBox 9", paras[0] if paras else "", label="paragraph 1")
        _opt_text(ctx, sl, "TextBox 11", paras[1] if len(paras) > 1 else "", label="paragraph 2")
    set_picture(ctx, sl, "Group 12", s.get("image"))


_AGENDA = [  # (number box, label box) in reading order 01..08
    ("TextBox 54", "TextBox 58"), ("TextBox 56", "TextBox 60"), ("TextBox 62", "TextBox 64"),
    ("TextBox 66", "TextBox 68"), ("TextBox 55", "TextBox 59"), ("TextBox 57", "TextBox 61"),
    ("TextBox 63", "TextBox 65"), ("TextBox 67", "TextBox 69"),
]


def fill_agenda(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 53"), s.get("title") or "Agenda", avail_h=2.58)
    items = s.get("items") or []
    groups = [g for g in sl.shapes if g.shape_type == 6]
    for i, (num_name, label_name) in enumerate(_AGENDA):
        num, label = shape(sl, num_name), shape(sl, label_name)
        if i < len(items):
            set_line(ctx, num, f"{i + 1:02d}")
            set_line(ctx, label, items[i], floor=0.55, label="agenda item")
            continue
        # unused row: remove its number, label, pill and circle (found by position)
        cy = (num.top + num.height / 2) / EMU
        cx = num.left / EMU
        for g in groups:
            gx, gy, gw, gh = _inches(g)
            same_row = abs((gy + gh / 2) - cy) < 0.35
            same_col = (gx < 10) == (cx < 10)
            if same_row and same_col and gh < 1.3 and g._element.getparent() is not None:
                g._element.getparent().remove(g._element)
        remove(sl, num_name, label_name)
    set_picture(ctx, sl, "Group 70", s.get("image"))


def fill_section(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Section", avail_h=3.0)
    _opt_text(ctx, sl, "TextBox 8", s.get("text"))
    set_picture(ctx, sl, "Group 6", s.get("image"))


def _fill_cards(ctx, sl, cards, slots):
    """slots: [(group, header box, body box)]; unused cards are removed."""
    for i, (grp, head, body) in enumerate(slots):
        if i < len(cards):
            set_line(ctx, shape(sl, head), cards[i].get("header", ""), floor=0.6, glyph=0.52, label="card header")
            set_text(ctx, shape(sl, body), cards[i].get("body", ""), label="card body")
        else:
            remove(sl, grp, head, body)


def fill_three_cards(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Key Points", avail_h=5.3)
    _opt_text(ctx, sl, "TextBox 23", s.get("text"))
    _fill_cards(ctx, sl, s.get("cards") or [], [
        ("Group 14", "TextBox 19", "TextBox 17"), ("Group 11", "TextBox 20", "TextBox 18"),
        ("Group 8", "TextBox 22", "TextBox 21")])
    set_picture(ctx, sl, "Group 6", s.get("image"))


def fill_two_cards(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 18"), s.get("title") or "Methodology", avail_h=2.0)
    _opt_text(ctx, sl, "TextBox 14", s.get("text"), capacity=480)
    _fill_cards(ctx, sl, s.get("cards") or [], [
        ("Group 9", "TextBox 13", "TextBox 12"), ("Group 4", "TextBox 8", "TextBox 7")])
    set_picture(ctx, sl, "Group 2", s.get("image"))


def fill_team(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 31"), s.get("title") or "Framework", avail_h=1.55)
    _opt_text(ctx, sl, "TextBox 32", s.get("text"))
    items = s.get("overview_items") or []
    if items:
        set_line(ctx, shape(sl, "TextBox 37"), s.get("overview_title") or "Overview")
        set_text(ctx, shape(sl, "TextBox 36"), items, capacity=110, label="overview_items")
    else:
        remove(sl, "TextBox 36", "TextBox 37", "Group 33")
    slots = [  # photo group, white label card, name, role, note   (left to right)
        ("Group 4", "Group 11", "TextBox 27", "TextBox 21", "TextBox 24"),
        ("Group 2", "Group 8", "TextBox 26", "TextBox 20", "TextBox 23"),
        ("Group 6", "Group 14", "TextBox 25", "TextBox 19", "TextBox 22"),
    ]
    people = s.get("people") or []
    for i, (photo, card, name, role, note) in enumerate(slots):
        if i >= len(people):
            remove(sl, photo, card, name, role, note)
            continue
        p = people[i]
        set_line(ctx, shape(sl, name), p.get("name", ""), floor=0.55, glyph=0.52, label="person name")
        set_line(ctx, shape(sl, role), p.get("role", ""), floor=0.6, label="person role") if p.get("role") \
            else remove(sl, role)
        set_line(ctx, shape(sl, note), p.get("note", ""), floor=0.6) if p.get("note") else remove(sl, note)
        set_picture(ctx, sl, photo, p.get("image"), person=True, label="photo")
    set_picture(ctx, sl, "Group 17", s.get("image"))


def fill_people(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Participants", avail_h=3.9)
    _opt_text(ctx, sl, "TextBox 6", s.get("text"))
    _opt_text(ctx, sl, "TextBox 7", s.get("text2"))
    slots = [  # photo group, name, role   (top-left, top-right, bottom-left, bottom-right)
        ("Group 10", "TextBox 21", "TextBox 17"), ("Group 14", "TextBox 23", "TextBox 19"),
        ("Group 8", "TextBox 20", "TextBox 16"), ("Group 12", "TextBox 22", "TextBox 18"),
    ]
    people = s.get("people") or []
    for i, (photo, name, role) in enumerate(slots):
        if i >= len(people):
            remove(sl, photo, name, role)
            continue
        p = people[i]
        set_line(ctx, shape(sl, name), p.get("name", ""), floor=0.55, glyph=0.52, label="person name")
        set_line(ctx, shape(sl, role), p.get("role", ""), floor=0.6) if p.get("role") else remove(sl, role)
        set_picture(ctx, sl, photo, p.get("image"), person=True, label="photo")


def fill_chart(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Data", avail_h=2.8)
    _opt_text(ctx, sl, "TextBox 6", s.get("text"))
    remove(sl, "Picture 10")                       # decorative pictogram of the template
    if not add_chart(ctx, sl, "Picture 11", s.get("chart"), "column", dark=True, box=(1.6, 1.7, 7.3, 7.9)):
        ctx.warn("no chart data - the chart card is empty (add \"chart\")")
    stat = s.get("stat") or {}
    if stat.get("value"):
        _add_stat(sl, stat["value"], stat.get("label", ""))


def _add_stat(sl, value, label):
    """Big number + caption, laid out like the template's own "65%" block."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    value = _clean(value)[:8]
    size = 96 if len(value) <= 4 else 72 if len(value) <= 6 else 56
    box = sl.shapes.add_textbox(Inches(11.11), Inches(7.9), Inches(4.4), Inches(1.6))
    box.text_frame.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(box.text_frame, side, 0)
    para = box.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = value
    run.font.size, run.font.bold, run.font.name = Pt(size), True, "Open Sauce Bold"
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if label:
        cap = sl.shapes.add_textbox(Inches(15.6), Inches(8.1), Inches(3.3), Inches(1.6))
        cap.text_frame.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(cap.text_frame, side, 0)
        r2 = cap.text_frame.paragraphs[0].add_run()
        r2.text = _clean(label)[:120]
        r2.font.size, r2.font.name = Pt(18), "HK Grotesk Italics"
        r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def fill_timeline(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Timeline", avail_h=2.2)
    _opt_text(ctx, sl, "TextBox 37", s.get("text"))
    slots = [  # header pill, header text, top item, bottom item, divider on the LEFT of the column
        ("Group 9", "TextBox 25", "TextBox 15", "TextBox 16", None),
        ("Group 12", "TextBox 26", "TextBox 17", "TextBox 18", "AutoShape 34"),
        ("Group 19", "TextBox 27", "TextBox 29", "TextBox 30", "AutoShape 35"),
        ("Group 22", "TextBox 28", "TextBox 31", "TextBox 32", "AutoShape 36"),
    ]
    cols = s.get("columns") or []
    for i, (pill, head, top, bottom, divider) in enumerate(slots):
        if i >= len(cols):
            remove(sl, pill, head, top, bottom, *( [divider] if divider else []))
            continue
        c = cols[i]
        set_line(ctx, shape(sl, head), c.get("header", ""), floor=0.55, glyph=0.55, label="timeline header")
        items = c.get("items") or []
        for name, idx in ((top, 0), (bottom, 1)):
            if idx < len(items):
                set_text(ctx, shape(sl, name), items[idx], capacity=80, label="timeline item")
            else:
                remove(sl, name)


def fill_analysis(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 6"), s.get("title") or "Analysis", avail_h=2.6)
    _opt_text(ctx, sl, "TextBox 2", s.get("text"), capacity=470 if s.get("chart2") else 1000)
    if not add_chart(ctx, sl, "Picture 11", s.get("chart"), "line", dark=False, box=(10.6, 1.6, 7.9, 8.0)):
        ctx.warn("no chart data - the white card is empty (add \"chart\")")
    add_chart(ctx, sl, "Picture 10", s.get("chart2"), "bar", dark=True, box=(1.0, 6.9, 8.6, 3.5))


def fill_stat_chart(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 7"), s.get("title") or "Analysis", avail_h=2.6)
    _opt_text(ctx, sl, "TextBox 2", s.get("text"))
    if s.get("stat"):
        set_line(ctx, shape(sl, "TextBox 8"), s["stat"], floor=0.45, glyph=0.6, label="stat")
        _opt_text(ctx, sl, "TextBox 3", s.get("stat_text"), label="stat_text")
    else:
        remove(sl, "TextBox 8")
        if s.get("stat_text"):
            tb = shape(sl, "TextBox 3")
            tb.left, tb.width = shape(sl, "TextBox 2").left, shape(sl, "TextBox 2").width
            set_text(ctx, tb, s["stat_text"], capacity=500, label="stat_text")
        else:
            remove(sl, "TextBox 3")
    if not add_chart(ctx, sl, "Picture 12", s.get("chart"), "radar", dark=False, box=(10.6, 1.6, 7.9, 8.0)):
        ctx.warn("no chart data - the white card is empty (add \"chart\")")


def fill_gallery(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Gallery", avail_h=3.0)
    shape(sl, "TextBox 5").width = int(6.0 * EMU)     # the photos start at x = 7.4in
    _opt_text(ctx, sl, "TextBox 12", s.get("text"))
    images = s.get("images") or []
    for i, grp in enumerate(("Group 6", "Group 8", "Group 10")):
        set_picture(ctx, sl, grp, images[i] if i < len(images) else None)


def fill_testimonials(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Testimonial", avail_h=2.1)
    _opt_text(ctx, sl, "TextBox 9", s.get("lead"), label="lead")
    slots = [  # photo, name, text, first star group number   (TL, TR, BL, BR)
        ("Group 14", "TextBox 80", "TextBox 84", 48), ("Group 16", "TextBox 81", "TextBox 85", 63),
        ("Group 10", "TextBox 78", "TextBox 82", 18), ("Group 12", "TextBox 79", "TextBox 83", 33),
    ]
    items = s.get("items") or []
    for i, (photo, name, text, star0) in enumerate(slots):
        stars = [f"Group {star0 + 3 * k}" for k in range(5)]
        if i >= len(items):
            remove(sl, photo, name, text, *stars)
            continue
        it = items[i]
        set_line(ctx, shape(sl, name), it.get("name", ""), floor=0.6, glyph=0.52, label="name")
        set_text(ctx, shape(sl, text), it.get("text", ""), label="testimonial text")
        for k, star in enumerate(stars):
            try:
                clr = shape(sl, star)._element.xpath(".//a:solidFill/a:srgbClr")
            except KeyError:
                continue
            for c in clr:
                c.set("val", "FFFFFF" if k < it.get("rating", 5) else "5B5B5B")
        set_picture(ctx, sl, photo, it.get("image"), person=True, label="photo")


def fill_closing(ctx, sl, s):
    set_title(ctx, shape(sl, "TextBox 5"), s.get("title") or "Thank You", avail_h=3.6, floor=0.3, glyph=0.54)
    if s.get("subtitle"):
        set_text(ctx, shape(sl, "TextBox 10"), s["subtitle"], capacity=70, label="subtitle")
    else:
        remove(sl, "TextBox 10")
    if s.get("tagline"):
        set_line(ctx, shape(sl, "TextBox 9"), s["tagline"], label="tagline")
    else:
        remove(sl, "TextBox 9", "Group 6")
    if s.get("organization"):
        set_line(ctx, shape(sl, "TextBox 11"), s["organization"], label="organization")
    else:
        remove(sl, "TextBox 11")


FILLERS = {
    "cover": fill_cover, "intro": fill_intro, "agenda": fill_agenda, "section": fill_section,
    "three_cards": fill_three_cards, "team": fill_team, "two_cards": fill_two_cards,
    "people": fill_people, "chart": fill_chart, "timeline": fill_timeline, "analysis": fill_analysis,
    "stat_chart": fill_stat_chart, "gallery": fill_gallery, "testimonials": fill_testimonials,
    "closing": fill_closing,
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def _result(ok, **kw):
    out = {"ok": ok}
    out.update(kw)
    return out


def build_deck(spec, filename: str = "research_deck.pptx", output_dir=None, image_dirs=None,
               template=None, keep_order: bool = False, auto_agenda: bool = True, strict: bool = False) -> dict:
    """Build the deck.  Returns a result dict and never raises.

    {"ok": True,  "path", "slides", "order": [layouts], "warnings": [...]}
    {"ok": False, "error", "hint", "warnings": [...]}
    """
    ctx = Ctx(image_dirs)
    try:
        from pptx import Presentation
    except ImportError:
        return _result(False, error="Python package `python-pptx` is not installed.",
                       hint="Run: pip install python-pptx", warnings=[])

    try:
        slides, meta = parse_slides(spec)
    except ValueError as e:
        return _result(False, error=f"Invalid JSON: {e}",
                       hint='Send {"slides":[{"layout":"cover","title":"..."}, ...]}', warnings=[])
    if meta.get("truncated"):
        ctx.warn("JSON was cut off (token limit?) - only the complete slides were used")
    slides, norm_warnings = normalize_spec(slides, meta, keep_order=keep_order, auto_agenda=auto_agenda)
    ctx.warnings += norm_warnings
    if not slides:
        return _result(False, error="No usable slides in the spec.",
                       hint='Each slide is an object such as {"layout":"section","title":"...","text":"..."}',
                       warnings=ctx.warnings)
    if strict and ctx.warnings:
        return _result(False, error="Spec needed repairs (strict mode).", hint="Fix the listed items.",
                       warnings=ctx.warnings)

    tpl = Path(template or os.environ.get(TEMPLATE_ENV) or DEFAULT_TEMPLATE)
    if not tpl.is_file():
        return _result(False, error=f"Template not found: {tpl}",
                       hint="Put your copy of the template at assets/template.pptx "
                            "(see references/template_setup.md) or pass --template.", warnings=ctx.warnings)

    name = str(filename or "research_deck.pptx")
    if not name.lower().endswith(".pptx"):
        name += ".pptx"
    out_path = Path(name).expanduser()
    if not out_path.is_absolute():
        out_path = Path(output_dir or os.environ.get(OUTPUT_DIR_ENV) or (Path.cwd() / "output")) / out_path
    out_path = out_path.resolve()

    try:
        prs = Presentation(str(tpl))
        originals = list(prs.slides)
        if len(originals) < 15:
            return _result(False, error=f"Template has {len(originals)} slides, 15 expected.",
                           hint="Use the unmodified template (see references/template_setup.md).",
                           warnings=ctx.warnings)
        for i, s in enumerate(slides, 1):
            ctx.tag = f"slide {i} ({s['layout']})"
            new = duplicate_slide(prs, originals[TEMPLATE_SLIDE[s["layout"]] - 1])
            try:
                FILLERS[s["layout"]](ctx, new, s)
            except KeyError as e:
                return _result(False, error=f"{ctx.tag}: {e.args[0] if e.args else e}",
                               hint="Restore assets/template.pptx from the original file.", warnings=ctx.warnings)
            if s.get("notes"):
                new.notes_slide.notes_text_frame.text = _clean(s["notes"])
        ctx.tag = ""
        delete_slides(prs, originals)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
    except PermissionError:
        return _result(False, error=f"Cannot write {out_path} (is the file open in PowerPoint?)",
                       hint="Close the file or choose another name.", warnings=ctx.warnings)
    except Exception as e:
        return _result(False, error=f"Build failed: {type(e).__name__}: {e}", hint="", warnings=ctx.warnings)

    return _result(True, path=str(out_path), slides=len(slides),
                   order=[s["layout"] for s in slides], warnings=ctx.warnings)


def format_result(res: dict) -> str:
    if res.get("ok"):
        lines = [f"OK: created {res['path']} ({res['slides']} slides)",
                 "ORDER: " + " > ".join(res.get("order", []))]
    else:
        lines = [f"ERROR: {res.get('error')}"]
        if res.get("hint"):
            lines.append(f"HINT: {res['hint']}")
    ws = res.get("warnings") or []
    if ws:
        lines.append(f"AUTO-FIXED / WARNINGS ({len(ws)}):")
        lines += [f"- {w}" for w in ws[:14]]
        if len(ws) > 14:
            lines.append(f"- ... and {len(ws) - 14} more")
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build a research-proposal deck from a JSON spec (see SKILL.md).")
    ap.add_argument("spec", help="JSON file, or '-' for stdin")
    ap.add_argument("-o", "--output", default="research_deck.pptx")
    ap.add_argument("--template", help="Template .pptx (default assets/template.pptx)")
    ap.add_argument("--keep-order", action="store_true", help="Do not re-order slides to the standard section order")
    ap.add_argument("--no-agenda", action="store_true", help="Do not add an agenda slide automatically")
    ap.add_argument("--strict", action="store_true", help="Fail instead of repairing the spec")
    ap.add_argument("--json", action="store_true", help="Print the result as JSON")
    args = ap.parse_args(argv)

    if args.spec == "-":
        raw, image_dirs = sys.stdin.read(), []
    else:
        p = Path(args.spec)
        if not p.is_file():
            print(f"ERROR: spec file not found: {p}")
            return 2
        raw, image_dirs = p.read_text(encoding="utf-8-sig"), [p.resolve().parent]
    out = Path(args.output)
    res = build_deck(raw, filename=str(out if out.is_absolute() else Path.cwd() / out), image_dirs=image_dirs,
                     template=args.template, keep_order=args.keep_order, auto_agenda=not args.no_agenda,
                     strict=args.strict)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    print(json.dumps(res, ensure_ascii=False, indent=2) if args.json else format_result(res))
    return 0 if res.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
